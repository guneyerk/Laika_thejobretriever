#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Laika, The Job Retriever - helper CLI.

Handles the bookkeeping for the Laika job-search agent: the master-folder
layout, profile change detection, LinkedIn fetching, de-duplication, the Excel
tracker, role cards and feedback statistics. The judgement work (rating roles
against the CV, explaining fit, proposing criteria changes) is done by the
agent itself.

Run `python laika.py --help` for the command list.
"""
from __future__ import annotations

import argparse
import datetime as dt
import hashlib
import html as htmllib
import json
import os
import random
import re
import shutil
import sys
import time
import unicodedata
import urllib.error
import urllib.parse
import urllib.request
import zipfile
from pathlib import Path

try:  # Windows consoles default to cp1252; job data is full of Turkish/other characters.
    sys.stdout.reconfigure(encoding="utf-8")
    sys.stderr.reconfigure(encoding="utf-8")
except Exception:
    pass

VERSION = "1.2.0"
GLOBAL_CONFIG = Path.home() / ".laika" / "config.json"
PROFILE_FILE = ".laika-profile.json"   # put this file in a folder to give that folder its own Laika profile


def resolve_config(workspace: str | None = None) -> tuple[Path, str]:
    """Which profile pointer to use: $LAIKA_CONFIG, else the nearest folder profile, else the global one."""
    env = os.environ.get("LAIKA_CONFIG")
    if env:
        return Path(env), "LAIKA_CONFIG"
    start = Path(workspace) if workspace and Path(workspace).is_dir() else Path.cwd()
    start = start.resolve()
    for d in (start, *start.parents):
        f = d / PROFILE_FILE
        if f.is_file():
            return f, "folder"
    return GLOBAL_CONFIG, "global"


CONFIG_PATH, CONFIG_KIND = resolve_config()

# ----------------------------------------------------------------------------
# Master folder layout
# ----------------------------------------------------------------------------
DIR_CV = "CV"
DIR_CERTS = "Licenses and Certifications"   # optional; read together with the CV
DIR_TAILORED = "Tailored CVs"               # one folder per role: tailored CV (.docx + .pdf) and its content.json
DIR_LETTERS = "Cover Letters"               # one folder per role: cover letter (.docx + .pdf) and its letter.json
DIR_STYLE = "Writing Style"                 # the user's writing style for cover letters (+ Samples/)
DIR_LOCATIONS = "Locations"
DIR_ROLES = "Roles"
DIR_CRITERIA = "Search Criteria"
DIR_SITES = "Job Sites"
DIR_SCANNED = "Roles Scanned"
DIR_APPLIED = "Roles Applied"
DIR_IGNORED = "Roles Ignored"
DIR_INTERNAL = ".laika"
ALL_DIRS = [DIR_CV, DIR_CERTS, DIR_LOCATIONS, DIR_ROLES, DIR_CRITERIA, DIR_SITES,
            DIR_SCANNED, DIR_APPLIED, DIR_IGNORED, DIR_TAILORED, DIR_LETTERS, DIR_INTERNAL]

F_LOCATIONS = f"{DIR_LOCATIONS}/locations.txt"
F_ROLES = f"{DIR_ROLES}/roles.txt"
F_CRITERIA = f"{DIR_CRITERIA}/criteria.txt"
F_SITES = f"{DIR_SITES}/sites.txt"
F_CHANGELOG = f"{DIR_CRITERIA}/changelog.md"
F_CV_SUMMARY = f"{DIR_CV}/Laika CV Summary.md"
F_CV_BASE = f"{DIR_TAILORED}/Base CV content.json"   # the CV as structured content, the start of every tailored CV
F_STYLE = f"{DIR_STYLE}/Laika Writing Style.md"      # how the user writes; Laika keeps it up to date from feedback
DIR_STYLE_SAMPLES = f"{DIR_STYLE}/Samples"
F_TRACKER = "Job Tracker.xlsx"
F_DB = f"{DIR_INTERNAL}/jobs.json"
F_STATE = f"{DIR_INTERNAL}/state.json"

STATUSES = ["Pending", "Applied", "Ignored", "Low Match"]
STATUS_ORDER = {s: i for i, s in enumerate(STATUSES)}
CARD_DIRS = {"Applied": DIR_APPLIED, "Ignored": DIR_IGNORED}

LIST_SPECS = {
    "locations": (F_LOCATIONS, "Locations", [
        "One location per line, e.g. 'Istanbul, Turkey' or 'London, United Kingdom'.",
        "Use 'Remote' for remote roles anywhere, or 'Remote (Turkey)' for remote roles in one country.",
    ]),
    "roles": (F_ROLES, "Roles", [
        "One job title per line. Laika runs a separate search for each one.",
        "Add local-language titles too if you want them searched everywhere (e.g. 'Vertriebsleiter').",
    ]),
    "sites": (F_SITES, "Job Sites", [
        "One site per line:   Name | URL | method | search titles (optional)",
        "method: api     = Laika's built-in LinkedIn fetcher",
        "        browser = searched in a browser (Claude desktop's built-in browser or Claude in Chrome)",
        "        web     = found through web search and page fetches",
        "search titles: optional '; '-separated keywords to search on THIS site instead of roles.txt,",
        "        e.g. local-language titles:  kariyer.net | https://www.kariyer.net | browser | Satis Direktoru; Pazarlama Direktoru",
    ]),
}
DEFAULT_SITES = [
    "LinkedIn | https://www.linkedin.com/jobs | api",
]
# Regional boards with built-in search URLs (see the search playbook). Onboarding offers them when
# the user's locations match; any other board works too with method "browser" or "web".
KNOWN_SITES = {
    "turkey": "kariyer.net | https://www.kariyer.net | browser",
    "uae": "GulfTalent | https://www.gulftalent.com | browser",
}

CRITERIA_SPEC = [
    ("seniority", "", "Target seniority, e.g. Mid, Senior, Lead, Manager, Head/Director, Executive"),
    ("team_leadership", "any", "required | preferred | any | not wanted"),
    ("work_mode", "any", "any, or a list of: onsite, hybrid, remote"),
    ("employment_type", "full-time", "full-time, part-time, contract, freelance, internship"),
    ("languages", "", "Working languages you are comfortable with"),
    ("industries_preferred", "", "Industries / domains you would like (raises the score)"),
    ("industries_excluded", "", "Industries / domains to avoid"),
    ("exclude_keywords", "", "Roles whose TITLE contains any of these words are auto-marked Low Match"),
    ("companies_excluded", "", "Companies to skip (current employer, recruiters, ...)"),
    ("salary_expectation", "", "Optional, free text"),
    ("other_preferences", "", "Anything else Laika should weigh (free text)"),
    ("min_relevancy_score", "60", "Roles scoring below this are registered as Low Match and not shown"),
    ("max_results_per_query", "25", "How many LinkedIn results to pull per role x location"),
    ("first_search_lookback_days", "14", "How far back the first search goes (also used for a newly added role / location / site)"),
]
CRITERIA_DEFAULTS = {k: v for k, v, _ in CRITERIA_SPEC}

UA = ("Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
      "(KHTML, like Gecko) Chrome/129.0.0.0 Safari/537.36")

STOPWORDS = {"and", "or", "the", "of", "for", "a", "an", "in", "at", "to", "with", "on",
             "ve", "ile", "sr", "jr", "m", "f", "d", "w", "x", "h"}
COMPANY_NOISE = {"a", "s", "as", "a.s", "anonim", "sirketi", "sti", "ltd", "limited", "inc", "llc",
                 "gmbh", "plc", "corp", "corporation", "co", "company", "tic", "ticaret", "san",
                 "sanayi", "ve", "bv", "ag", "sa", "srl", "turkiye", "turkey", "group", "grubu",
                 "holding", "the"}
TR_MAP = str.maketrans({"ı": "i", "İ": "i", "ş": "s", "Ş": "s", "ğ": "g", "Ğ": "g",
                        "ç": "c", "Ç": "c", "ö": "o", "Ö": "o", "ü": "u", "Ü": "u"})


# ----------------------------------------------------------------------------
# Small helpers
# ----------------------------------------------------------------------------
def now_local() -> dt.datetime:
    return dt.datetime.now().replace(microsecond=0)


def today_str() -> str:
    return dt.date.today().isoformat()


def die(msg: str, code: int = 1):
    print(f"ERROR: {msg}")
    sys.exit(code)


def read_json(path: Path, default):
    if not path.exists():
        return default
    with open(path, encoding="utf-8") as fh:
        return json.load(fh)


def write_json(path: Path, data) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_name(path.name + ".tmp")
    with open(tmp, "w", encoding="utf-8") as fh:
        json.dump(data, fh, ensure_ascii=False, indent=2)
    os.replace(tmp, path)


def norm(s: str) -> str:
    s = (s or "").translate(TR_MAP)
    s = unicodedata.normalize("NFKD", s)
    s = "".join(c for c in s if not unicodedata.combining(c)).lower()
    return re.sub(r"[^a-z0-9]+", " ", s).strip()


def tokens(s: str) -> list[str]:
    return [t for t in norm(s).split() if t not in STOPWORDS]


def norm_company(s: str) -> str:
    return " ".join(t for t in norm(s).split() if t not in COMPANY_NOISE) or norm(s)


def norm_title(s: str) -> str:
    s = re.sub(r"\((?:[mwfdhx]\s*/\s*)+[mwfdhx]\)", " ", s or "", flags=re.I)  # (m/f/d) etc.
    return norm(s)


LOC_NOISE = {"greater", "metropolitan", "area", "region", "bolgesi", "ili", "merkez", "city"}


def norm_city(s: str) -> str:
    """Best-effort city key: 'Kartal, Istanbul, Turkiye', 'Greater Istanbul' and 'Istanbul(Avr.)' -> 'istanbul'."""
    n = norm(s)
    if not n:
        return ""
    if n.startswith(("remote", "uzaktan")) or " remote" in f" {n}":
        return "remote"
    s = re.sub(r"\(.*?\)|\+\s*\d+.*$", " ", s)  # '(Avr.)', '(Tumu) +80 il daha'
    parts = [norm(p) for p in re.split(r"[,/]", s) if norm(p)]
    if not parts:
        return n.split()[0]
    if len(parts) >= 3 and parts[-1] in ("turkiye", "turkey"):  # LinkedIn TR: district, province, country
        cand = parts[-2]
    else:  # city, region, country | city, country | city
        cand = parts[0]
    return " ".join(t for t in cand.split() if t not in LOC_NOISE) or cand


def fingerprint(company: str, title: str, location: str) -> str:
    return f"{norm_company(company)}|{norm_title(title)}|{norm_city(location)}"


KEEP_QUERY_KEYS = {"jk", "jobid", "job_id", "id", "gh_jid", "currentjobid", "vjk", "jobkey", "postingid"}


def job_key(url: str) -> str:
    u = (url or "").strip()
    m = re.search(r"linkedin\.com/.*?(?:jobs/view/(?:[^/?#]*?-)?(\d{6,})|currentJobId=(\d{6,}))", u)
    if m:
        return "linkedin:" + (m.group(1) or m.group(2))
    m = re.search(r"kariyer\.net/is-ilani/[^?#]*?-(\d{5,})(?:[/?#]|$)", u)
    if m:
        return "kariyer:" + m.group(1)
    p = urllib.parse.urlsplit(u)
    host = p.netloc.lower()
    host = host[4:] if host.startswith("www.") else host
    q = [(k, v) for k, v in urllib.parse.parse_qsl(p.query) if k.lower() in KEEP_QUERY_KEYS]
    qs = ("?" + urllib.parse.urlencode(sorted(q))) if q else ""
    return f"url:{host}{p.path.rstrip('/')}{qs}"


def html_inline(s: str) -> str:
    s = re.sub(r"<[^>]+>", " ", s or "")
    return re.sub(r"\s+", " ", htmllib.unescape(s)).strip()


def html_to_text(s: str) -> str:
    s = re.sub(r"(?is)<(script|style)[^>]*>.*?</\1>", "", s or "")
    s = re.sub(r"(?i)<br\s*/?>", "\n", s)
    s = re.sub(r"(?i)<li[^>]*>", "\n- ", s)
    s = re.sub(r"(?i)</(p|div|ul|ol|h\d|tr)>", "\n", s)
    s = re.sub(r"<[^>]+>", "", s)
    s = htmllib.unescape(s)
    lines = [re.sub(r"[ \t ]+", " ", ln).strip() for ln in s.splitlines()]
    out, blank = [], False
    for ln in lines:
        if not ln:
            if not blank and out:
                out.append("")
            blank = True
        else:
            out.append(ln)
            blank = False
    return "\n".join(out).strip()


def safe_name(s: str, maxlen: int = 70) -> str:
    s = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "", s or "").strip()
    s = re.sub(r"\s+", " ", s).rstrip(". ")
    return s[:maxlen].rstrip(". ") or "untitled"


def split_list(value: str) -> list[str]:
    return [p.strip() for p in re.split(r"[,;\n]", value or "") if p.strip()]


def trim(s: str, n: int) -> str:
    s = s or ""
    return s if len(s) <= n else s[: n - 1].rstrip() + "…"


REQ_RE = re.compile(r"(requirement|qualification|what you.ll bring|what we.re looking for|who you are|about you|"
                    r"you have|you bring|experience|profile|nitelik|anforderung|profil)", re.I)
WHERE_RE = re.compile(r"(remote|hybrid|on-?site|based in|location|relocat|work from|uzaktan)", re.I)


def description_excerpt(desc: str) -> str:
    """Compact view for scoring long runs: opening, where/remote mentions, and the requirements part."""
    head = desc[:550]
    m = REQ_RE.search(desc, 400)
    req = desc[m.start(): m.start() + 850] if m else desc[550:1300]
    where = sorted({desc[max(0, x.start() - 60): x.start() + 90].replace("\n", " ") for x in WHERE_RE.finditer(desc)})[:3]
    parts = [head]
    if where:
        parts.append("WHERE/REMOTE: " + " ... ".join(where))
    if req:
        parts.append("[..] " + req)
    return "\n".join(parts)


# ----------------------------------------------------------------------------
# Config, master folder, state, db
# ----------------------------------------------------------------------------
def load_config() -> dict:
    return read_json(CONFIG_PATH, {})


def save_config(cfg: dict) -> None:
    write_json(CONFIG_PATH, cfg)


def get_master(required: bool = True) -> Path | None:
    p = os.environ.get("LAIKA_HOME") or load_config().get("master")
    if p and (Path(p) / F_STATE).exists():
        return Path(p)
    if required:
        print("STATE: NOT_INITIALIZED")
        if p:
            print(f"The configured master folder was not found or is not set up: {p}")
        print("Run onboarding, then: laika.py init --master \"<folder>\"")
        sys.exit(3)
    return None


def load_state(m: Path) -> dict:
    return read_json(m / F_STATE, {})


def save_state(m: Path, st: dict) -> None:
    write_json(m / F_STATE, st)


def load_db(m: Path) -> dict:
    return read_json(m / F_DB, {"version": 1, "next_id": 1, "jobs": []})


def save_db(m: Path, db: dict) -> None:
    write_json(m / F_DB, db)


def new_job_id(db: dict) -> str:
    i = db.get("next_id", 1)
    db["next_id"] = i + 1
    return f"J{i:04d}"


def db_index(db: dict) -> tuple[dict, dict]:
    keys, fps = {}, {}
    for j in db["jobs"]:
        keys[j["key"]] = j
        if j.get("fingerprint"):
            fps[j["fingerprint"]] = j
    return keys, fps


def find_job(db: dict, ref: str) -> dict | None:
    ref = ref.strip().upper()
    for j in db["jobs"]:
        if j["id"] == ref:
            return j
    return None


# ----------------------------------------------------------------------------
# Profile files (txt lists, criteria, sites)
# ----------------------------------------------------------------------------
def read_lines(path: Path) -> list[str]:
    if not path.exists():
        return []
    out = []
    for line in path.read_text(encoding="utf-8-sig").splitlines():
        s = line.strip()
        if s and not s.startswith("#"):
            out.append(s)
    return out


def write_list(m: Path, kind: str, items: list[str]) -> None:
    rel, title, notes = LIST_SPECS[kind]
    seen, clean = set(), []
    for it in items:
        it = it.strip()
        if it and norm(it) not in seen:
            seen.add(norm(it))
            clean.append(it)
    header = [f"# Laika - {title}"] + [f"# {n}" for n in notes] + ["# Lines starting with # are ignored.", ""]
    (m / rel).parent.mkdir(parents=True, exist_ok=True)
    (m / rel).write_text("\n".join(header + clean) + "\n", encoding="utf-8")


def read_sites(m: Path) -> list[dict]:
    sites = []
    for line in read_lines(m / F_SITES):
        parts = [p.strip() for p in line.split("|")]
        name = parts[0]
        url = parts[1] if len(parts) > 1 else ""
        method = (parts[2].lower() if len(parts) > 2 else "")
        if method not in ("api", "browser", "web"):
            method = "api" if "linkedin" in (name + url).lower() else "web"
        keywords = split_list(parts[3].replace(";", ",")) if len(parts) > 3 else []
        sites.append({"name": name, "url": url, "method": method, "keywords": keywords})
    return sites


def match_terms(m: Path) -> list[str]:
    """Titles used for title matching: roles.txt plus every site's own search titles."""
    terms = list(read_lines(m / F_ROLES))
    for site in read_sites(m):
        terms += site.get("keywords", [])
    return terms


def read_criteria(m: Path) -> dict:
    vals = dict(CRITERIA_DEFAULTS)
    path = m / F_CRITERIA
    if path.exists():
        for line in path.read_text(encoding="utf-8-sig").splitlines():
            s = line.strip()
            if not s or s.startswith("#") or ":" not in s:
                continue
            k, v = s.split(":", 1)
            vals[k.strip().lower().replace(" ", "_")] = v.strip()
    return vals


def write_criteria(m: Path, vals: dict) -> None:
    lines = ["# Laika - Search Criteria",
             "# Format:  key: value   (lists are comma separated). Lines starting with # are ignored.",
             "# Edit freely - Laika notices changes at the start of every run.", ""]
    known = set()
    for k, default, desc in CRITERIA_SPEC:
        known.add(k)
        lines.append(f"# {desc}")
        lines.append(f"{k}: {vals.get(k, default)}")
        lines.append("")
    extras = [k for k in vals if k not in known]
    if extras:
        lines.append("# Custom criteria")
        for k in extras:
            lines.append(f"{k}: {vals[k]}")
    (m / F_CRITERIA).parent.mkdir(parents=True, exist_ok=True)
    (m / F_CRITERIA).write_text("\n".join(lines).rstrip() + "\n", encoding="utf-8")


def crit_int(crit: dict, key: str) -> int:
    try:
        return int(str(crit.get(key, CRITERIA_DEFAULTS[key])).strip())
    except (ValueError, KeyError):
        return int(CRITERIA_DEFAULTS[key])


def log_change(m: Path, what: str, reason: str = "") -> None:
    path = m / F_CHANGELOG
    if not path.exists():
        path.write_text("# Search criteria change log\n\n", encoding="utf-8")
    line = f"- {now_local():%Y-%m-%d %H:%M} - {what}" + (f"  (reason: {reason})" if reason else "") + "\n"
    with open(path, "a", encoding="utf-8") as fh:
        fh.write(line)


# ----------------------------------------------------------------------------
# Change detection
# ----------------------------------------------------------------------------
TEXT_EXT = {".txt", ".md"}
IMAGE_EXT = {".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp", ".tif", ".tiff", ".heic"}


def cv_files(m: Path) -> list[Path]:
    d = m / DIR_CV
    if not d.exists():
        return []
    return sorted(p for p in d.iterdir()
                  if p.is_file() and not p.name.startswith(("~$", ".")) and p.name != Path(F_CV_SUMMARY).name)


def cert_files(m: Path) -> list[Path]:
    """Documents in the optional Licenses and Certifications folder (subfolders included)."""
    d = m / DIR_CERTS
    if not d.exists():
        return []
    return sorted(p for p in d.rglob("*")
                  if p.is_file() and not any(part.startswith(("~$", ".")) for part in p.relative_to(d).parts))


def cert_rel(m: Path, p: Path) -> str:
    return f"{DIR_CERTS}/{p.relative_to(m / DIR_CERTS).as_posix()}"


def tracked_files(m: Path) -> list[str]:
    rels = [F_LOCATIONS, F_ROLES, F_CRITERIA, F_SITES, F_CV_SUMMARY]
    rels += [f"{DIR_CV}/{p.name}" for p in cv_files(m)]
    rels += [cert_rel(m, p) for p in cert_files(m)]
    return rels


def sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def snap_copy(m: Path, rel: str) -> Path:
    return m / DIR_INTERNAL / "snapshots" / rel.replace("/", "__")


def meaningful_lines(path: Path) -> list[str]:
    if not path.exists():
        return []
    return [ln.strip() for ln in path.read_text(encoding="utf-8-sig", errors="replace").splitlines()
            if ln.strip() and not ln.strip().startswith("#")]


def detect_changes(m: Path, st: dict) -> list[dict]:
    snap = st.get("snapshot", {})
    current = {rel: sha(m / rel) for rel in tracked_files(m) if (m / rel).exists()}
    changes = []
    for rel, h in current.items():
        if rel not in snap:
            changes.append({"file": rel, "change": "added"})
        elif snap[rel] != h:
            c = {"file": rel, "change": "modified"}
            if Path(rel).suffix.lower() in TEXT_EXT:
                old = meaningful_lines(snap_copy(m, rel))
                new = meaningful_lines(m / rel)
                c["added_lines"] = [ln for ln in new if ln not in old]
                c["removed_lines"] = [ln for ln in old if ln not in new]
            changes.append(c)
    for rel in snap:
        if rel not in current:
            changes.append({"file": rel, "change": "removed"})
    return changes


def take_snapshot(m: Path, st: dict, only: list[str] | None = None) -> None:
    snap = dict(st.get("snapshot", {})) if only else {}
    for rel in (only or tracked_files(m)):
        p = m / rel
        cp = snap_copy(m, rel)
        if p.exists():
            snap[rel] = sha(p)
            if p.suffix.lower() in TEXT_EXT:
                cp.parent.mkdir(parents=True, exist_ok=True)
                shutil.copyfile(p, cp)
        else:
            snap.pop(rel, None)
            if cp.exists():
                cp.unlink()
    st["snapshot"] = snap
    st["snapshot_at"] = now_local().isoformat()


# ----------------------------------------------------------------------------
# CV text extraction
# ----------------------------------------------------------------------------
def extract_text(p: Path) -> str | None:
    ext = p.suffix.lower()
    if ext in (".txt", ".md"):
        return p.read_text(encoding="utf-8-sig", errors="replace")
    if ext == ".docx":
        with zipfile.ZipFile(p) as z:
            xml = z.read("word/document.xml").decode("utf-8", errors="replace")
        xml = re.sub(r"</w:p>", "\n", xml)
        xml = re.sub(r"<w:tab/>", "\t", xml)
        xml = re.sub(r"<w:br[^>]*/>", "\n", xml)
        return htmllib.unescape(re.sub(r"<[^>]+>", "", xml)).strip()
    if ext == ".pdf":
        try:
            import pypdf  # type: ignore
        except ImportError:
            return None
        reader = pypdf.PdfReader(str(p))
        return "\n".join((pg.extract_text() or "") for pg in reader.pages).strip()
    return None


# ----------------------------------------------------------------------------
# HTTP + LinkedIn (public guest endpoints, polite pace)
# ----------------------------------------------------------------------------
class FetchError(Exception):
    pass


def http_get(url: str, tries: int = 3) -> str:
    last = None
    for attempt in range(tries):
        req = urllib.request.Request(url, headers={
            "User-Agent": UA,
            "Accept": "text/html,application/xhtml+xml,*/*;q=0.8",
            "Accept-Language": "en-US,en;q=0.9",
        })
        try:
            with urllib.request.urlopen(req, timeout=25) as r:
                return r.read().decode("utf-8", errors="replace")
        except urllib.error.HTTPError as e:
            last = f"HTTP {e.code}"
            if e.code in (400, 404, 410):
                return ""
            if e.code in (429, 999) or e.code >= 500:
                time.sleep(8 * (attempt + 1))
                continue
            raise FetchError(last)
        except (urllib.error.URLError, TimeoutError, ConnectionError, OSError) as e:
            last = str(e)
            time.sleep(3)
    raise FetchError(last or "unknown network error")


def polite_sleep(delay: float) -> None:
    time.sleep(random.uniform(delay * 0.7, delay * 1.3))


def parse_location(loc: str) -> tuple[str, bool]:
    """Return (search location, remote_only)."""
    n = norm(loc)
    if n.startswith(("remote", "uzaktan")):
        m = re.search(r"\(([^)]+)\)", loc) or re.search(r"[-,:]\s*(.+)$", loc)
        return (m.group(1).strip() if m else "Worldwide"), True
    return loc, False


WORK_MODE_CODES = {"onsite": "1", "on-site": "1", "office": "1", "remote": "2", "hybrid": "3"}


def work_mode_codes(crit: dict) -> list[str]:
    modes = [norm(x).replace(" ", "") for x in split_list(crit.get("work_mode", ""))]
    if not modes or "any" in modes:
        return []
    codes = sorted({WORK_MODE_CODES[x] for x in modes if x in WORK_MODE_CODES})
    return [] if len(codes) == 3 else codes


def linkedin_params(q: dict, start: int | None = None) -> dict:
    p = {"keywords": q.get("keywords") or q["role"], "location": q["search_location"], "f_TPR": f"r{q['seconds']}"}
    if q.get("work_types"):
        p["f_WT"] = ",".join(q["work_types"])
    if start is not None:
        p["start"] = str(start)
    return p


def parse_linkedin_cards(page: str) -> list[dict]:
    out = []
    for chunk in re.split(r"<li[\s>]", page):
        m = re.search(r"urn:li:jobPosting:(\d+)", chunk) or re.search(r"/jobs/view/[^\"?]*?-(\d{6,})", chunk)
        if not m:
            continue
        jid = m.group(1)

        def grab(pat):
            mm = re.search(pat, chunk, re.S)
            return html_inline(mm.group(1)) if mm else ""

        posted = re.search(r"<time[^>]*datetime=\"([\d-]+)\"", chunk)
        out.append({
            "title": grab(r"<h3[^>]*base-search-card__title[^>]*>(.*?)</h3>"),
            "company": grab(r"<h4[^>]*base-search-card__subtitle[^>]*>(.*?)</h4>"),
            "location": grab(r"<span[^>]*job-search-card__location[^>]*>(.*?)</span>"),
            "posted": posted.group(1) if posted else "",
            "url": f"https://www.linkedin.com/jobs/view/{jid}/",
        })
    return out


def parse_linkedin_detail(page: str) -> dict:
    d: dict = {}
    i = page.find("show-more-less-html__markup")
    if i >= 0:
        start = page.find(">", i) + 1
        end = page.find("<button", start)
        d["description"] = trim(html_to_text(page[start:end if end > 0 else start + 20000]), 6000)
    for k, v in re.findall(r"description__job-criteria-subheader\">\s*(.*?)\s*</h3>\s*<span[^>]*>\s*(.*?)\s*</span>",
                           page, re.S):
        d[norm(html_inline(k)).replace(" ", "_")] = html_inline(v)
    m = re.search(r"num-applicants__caption[^>]*>(.*?)</", page, re.S)
    if m:
        d["applicants"] = html_inline(m.group(1))
    m = re.search(r"posted-time-ago__text[^>]*>(.*?)</span>", page, re.S)
    if m:
        d["posted_ago"] = html_inline(m.group(1))
    m = re.search(r"compensation__salary[^>]*>(.*?)</div>", page, re.S)
    if m:
        d["salary"] = html_inline(m.group(1))
    if "No longer accepting applications" in page:
        d["closed"] = True
    return d


# ----------------------------------------------------------------------------
# Runs (one search = one run)
# ----------------------------------------------------------------------------
def run_path(m: Path, rid: str) -> Path:
    return m / DIR_INTERNAL / "runs" / rid / "run.json"


def load_run(m: Path, st: dict, required: bool = True) -> dict | None:
    rid = st.get("current_run")
    if rid and run_path(m, rid).exists():
        run = read_json(run_path(m, rid), None)
        if run and run.get("abandoned"):
            # A discarded run must stay discarded, even if another session re-linked it in state.json.
            st["current_run"] = None
            save_state(m, st)
            if required:
                die(f"Run {rid} was discarded on {run['abandoned_at'][:16].replace('T', ' ')} and can't be used "
                    "again. Do not re-link it; start a new search with `run start`.")
            return None
        return run
    if required:
        die("No search run is open. Start one with: laika.py run start")
    return None


def save_run(m: Path, run: dict) -> None:
    write_json(run_path(m, run["id"]), run)


def combo_key(site: str, role: str, loc: str) -> str:
    return f"{norm(site)}|{norm(role)}|{norm(loc)}"


def compute_plan(m: Path, st: dict, crit: dict, now: dt.datetime) -> list[dict]:
    roles = read_lines(m / F_ROLES)
    locs = read_lines(m / F_LOCATIONS)
    sites = read_sites(m)
    lookback = crit_int(crit, "first_search_lookback_days")
    searched = st.get("searched", {})
    wt = work_mode_codes(crit)

    def since_for(site, role, loc):
        last = searched.get(combo_key(site, role, loc))
        if last:
            return dt.datetime.fromisoformat(last).date(), False
        return (now - dt.timedelta(days=lookback)).date(), True

    plan = []
    n = 0
    for site in sites:
        for role in (site.get("keywords") or roles):
            if site["method"] == "browser":
                sinces = [since_for(site["name"], role, loc) for loc in locs]
                since = min(s for s, _ in sinces)
                n += 1
                q = {"qid": f"q{n:02d}", "site": site["name"], "method": "browser", "role": role,
                     "locations": locs, "since": since.isoformat(), "new_combo": any(f for _, f in sinces),
                     "url": site_search_url(site, role)}
                plan.append(q)
                continue
            for loc in locs:
                since, fresh = since_for(site["name"], role, loc)
                n += 1
                q = {"qid": f"q{n:02d}", "site": site["name"], "method": site["method"], "role": role,
                     "location": loc, "since": since.isoformat(), "new_combo": fresh}
                if site["method"] == "api":
                    search_loc, remote = parse_location(loc)
                    since_dt = dt.datetime.combine(since, dt.time.min)
                    # The public endpoint ignores the work-type filter (f_WT), so remote searches also
                    # add "remote" to the keywords; work mode is verified again while scoring.
                    q["keywords"] = f"{role} remote" if remote else role
                    q["search_location"] = search_loc
                    q["work_types"] = ["2"] if remote else wt
                    q["seconds"] = max(3600, int((now - since_dt).total_seconds()))
                    q["url"] = "https://www.linkedin.com/jobs/search/?" + urllib.parse.urlencode(linkedin_params(q))
                else:
                    domain = urllib.parse.urlsplit(site["url"]).netloc or site["url"]
                    q["url"] = site["url"]
                    q["search_hint"] = f'site:{domain} "{role}" {loc}'
                plan.append(q)
    return plan


def site_search_url(site: dict, role: str) -> str:
    url = site["url"].lower()
    if "kariyer.net" in url:
        return "https://www.kariyer.net/is-ilanlari?kw=" + urllib.parse.quote(role)
    if "gulftalent.com" in url:  # keyword search, UAE (country id from GulfTalent's own filter)
        return ("https://www.gulftalent.com/jobs/search?country=10111111000000&search_keyword="
                + urllib.parse.quote(role))
    return site["url"]


def window_text(run: dict) -> str:
    sinces = sorted({q["since"] for q in run["plan"]})
    if not sinces:
        return "-"
    if len(sinces) == 1:
        return f"since {sinces[0]}"
    return f"since {sinces[-1]} (newly added roles/locations/sites: since {sinces[0]})"


def run_min_since(run: dict, source: str | None = None) -> str:
    qs = [q for q in run["plan"] if not source or norm(q["site"]) == norm(source)] or run["plan"]
    return min((q["since"] for q in qs), default=today_str())


def title_match(title: str, roles: list[str]) -> float:
    tt = set(tokens(title))
    best = 0.0
    for r in roles:
        rt = set(tokens(r))
        if rt:
            best = max(best, len(rt & tt) / len(rt))
    return round(best, 2)


def auto_flags(item: dict, crit: dict) -> list[str]:
    flags = []
    nt = f" {norm(item.get('title', ''))} "
    for kw in split_list(crit.get("exclude_keywords", "")):
        nk = norm(kw)
        if nk and f" {nk} " in nt:
            flags.append(f"title contains excluded keyword '{kw}'")
    nc = norm_company(item.get("company", ""))
    for comp in split_list(crit.get("companies_excluded", "")):
        ncomp = norm_company(comp)
        if ncomp and nc and (ncomp == nc or f" {ncomp} " in f" {nc} "):
            flags.append(f"company '{comp}' is excluded")
    return flags


def add_candidate(run: dict, item: dict, idx: tuple[dict, dict], crit: dict, roles: list[str],
                  source: str, query: str = "", min_since: str | None = None) -> tuple[str, str]:
    """Add one job to the run. Returns (outcome, detail)."""
    title = (item.get("title") or "").strip()
    url = (item.get("url") or "").strip()
    if not title or not url:
        return "invalid", "missing title or url"
    company = (item.get("company") or "").strip()
    location = (item.get("location") or "").strip()
    posted = (item.get("posted") or "").strip()
    if not posted and item.get("age_days") is not None:
        try:
            posted = (dt.date.today() - dt.timedelta(days=int(item["age_days"]))).isoformat()
        except (TypeError, ValueError):
            posted = ""
    if posted and min_since and posted < min_since:
        return "old", f"posted {posted}, before {min_since}"
    key = job_key(url)
    fp = fingerprint(company or urllib.parse.urlsplit(url).netloc, title, location)
    keys, fps = idx
    if key in keys:
        return "dup_db", keys[key]["id"]
    if fp in fps:
        return "dup_db", fps[fp]["id"]
    for c in run["candidates"]:
        if c["key"] == key or c["fingerprint"] == fp:
            if item.get("description") and not (c.get("details") or {}).get("description"):
                c.setdefault("details", {})["description"] = trim(item["description"], 6000)
                return "merged", c["cid"]
            return "dup_run", c["cid"]
    cand = {
        "cid": f"c{len(run['candidates']) + 1:03d}",
        "key": key, "fingerprint": fp, "title": title, "company": company, "location": location,
        "url": url, "source": source, "posted": posted,
        "work_mode": (item.get("work_mode") or "").strip(),
        "query": query, "title_match": title_match(title, roles),
        "flags": auto_flags({"title": title, "company": company}, crit),
    }
    if item.get("description"):
        cand["details"] = {"description": trim(item["description"], 6000)}
    if item.get("notes"):
        cand["notes"] = item["notes"]
    run["candidates"].append(cand)
    return "new", cand["cid"]


# ----------------------------------------------------------------------------
# Excel tracker
# ----------------------------------------------------------------------------
TRACKER_COLUMNS = [
    ("Role", 42), ("Company Name", 30), ("Relevancy Score", 11), ("Search Date", 12), ("Status", 12),
    ("Location", 24), ("Link", 13), ("Source", 12), ("Posted", 12), ("Decision Date", 13),
    ("Why Relevant", 60), ("Possible Concerns", 50), ("My Notes", 30), ("Tailored CV", 12), ("Cover Letter", 12),
    ("Job ID", 9),
]


def match_status(value) -> str | None:
    if not isinstance(value, str):
        return None
    for s in STATUSES:
        if norm(value) == norm(s):
            return s
    return None


def excel_sync_in(m: Path, db: dict) -> list[tuple[dict, str, str]]:
    """Pick up Status / My Notes edits the user made directly in Excel."""
    path = m / F_TRACKER
    if not path.exists():
        return []
    try:
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception as e:  # noqa: BLE001 - unreadable workbook should not stop a run
        print(f"WARNING: could not read {F_TRACKER} to import manual edits ({e}).")
        return []
    changed = []
    try:
        ws = wb["Job Tracker"] if "Job Tracker" in wb.sheetnames else wb.active
        rows = ws.iter_rows(values_only=True)
        header = next(rows, None) or []
        col = {h: i for i, h in enumerate(header) if h}
        if "Job ID" not in col or "Status" not in col:
            return []
        by_id = {j["id"]: j for j in db["jobs"]}
        for r in rows:
            jid = r[col["Job ID"]] if len(r) > col["Job ID"] else None
            j = by_id.get(jid)
            if not j:
                continue
            new = match_status(r[col["Status"]])
            if new and new != j.get("excel_status") and new != j["status"]:
                old = j["status"]
                j["status"] = new
                j["status_date"] = today_str()
                j["status_via"] = "excel"
                changed.append((j, old, new))
            if "My Notes" in col and len(r) > col["My Notes"]:
                note = r[col["My Notes"]]
                j["user_notes"] = str(note).strip() if note else ""
    finally:
        wb.close()
    return changed


def excel_write(m: Path, db: dict, st: dict) -> bool:
    try:
        from openpyxl import Workbook
        from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.worksheet.datavalidation import DataValidation
        from openpyxl.utils import get_column_letter
    except ImportError:
        print("WARNING: openpyxl is not installed - run: python -m pip install openpyxl")
        st["excel_dirty"] = True
        return False

    wb = Workbook()
    ws = wb.active
    ws.title = "Job Tracker"
    head_fill = PatternFill("solid", fgColor="1F4E5F")
    head_font = Font(bold=True, color="FFFFFF")
    for i, (name, width) in enumerate(TRACKER_COLUMNS, start=1):
        c = ws.cell(row=1, column=i, value=name)
        c.fill, c.font = head_fill, head_font
        c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[1].height = 30

    def sort_key(j):
        return (-dt.date.fromisoformat(j["search_date"]).toordinal(), STATUS_ORDER.get(j["status"], 9),
                -(j["score"] if j.get("score") is not None else -1))

    jobs = sorted(db["jobs"], key=sort_key)
    link_font = Font(color="0563C1", underline="single")
    top = Alignment(vertical="top")
    wrap = Alignment(vertical="top", wrap_text=True)
    for r, j in enumerate(jobs, start=2):
        vals = [j["title"], j["company"], j.get("score"),
                dt.date.fromisoformat(j["search_date"]), j["status"], j.get("location", ""),
                "Open posting" if j.get("url") else "", j.get("source", ""),
                dt.date.fromisoformat(j["posted"]) if re.fullmatch(r"\d{4}-\d{2}-\d{2}", j.get("posted") or "") else (j.get("posted") or None),
                dt.date.fromisoformat(j["status_date"]) if j["status"] in ("Applied", "Ignored") and j.get("status_date") else None,
                j.get("why", ""), (j.get("concerns", "") + (f"\n[Your reason: {j['reason']}]" if j.get("reason") else "")).strip(),
                j.get("user_notes", ""), "Open CV" if j.get("tailored_cv") else "",
                "Open letter" if j.get("cover_letter") else "", j["id"]]
        for ci, v in enumerate(vals, start=1):
            c = ws.cell(row=r, column=ci, value=v)
            c.alignment = wrap if ci in (1, 2, 11, 12, 13) else top
            if isinstance(v, dt.date):
                c.number_format = "yyyy-mm-dd"
        if j.get("url"):
            lc = ws.cell(row=r, column=7)
            lc.hyperlink = j["url"]
            lc.font = link_font
        for col, key in ((14, "tailored_cv"), (15, "cover_letter")):
            if j.get(key):  # relative to the tracker, so the link survives moving the folder
                cc = ws.cell(row=r, column=col)
                cc.hyperlink = j[key]
                cc.font = link_font
    last = max(2, len(jobs) + 1)
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(TRACKER_COLUMNS))}{last}"
    dv = DataValidation(type="list", formula1='"' + ",".join(STATUSES) + '"', allow_blank=False,
                        showErrorMessage=True, errorTitle="Status", error="Pick one of: " + ", ".join(STATUSES))
    ws.add_data_validation(dv)
    dv.add(f"E2:E{last}")
    fills = {"Applied": "C6EFCE", "Ignored": "D9D9D9", "Pending": "FFEB9C", "Low Match": "F2F2F2"}
    for s, color in fills.items():
        ws.conditional_formatting.add(f"E2:E{last}", CellIsRule(operator="equal", formula=[f'"{s}"'],
                                                                 fill=PatternFill("solid", fgColor=color)))
    ws.conditional_formatting.add(f"C2:C{last}", ColorScaleRule(
        start_type="num", start_value=0, start_color="F8696B", mid_type="num", mid_value=60,
        mid_color="FFEB84", end_type="num", end_value=100, end_color="63BE7B"))

    log = wb.create_sheet("Search Log")
    log_cols = [("Run", 18), ("Search Date", 12), ("Window Start", 13), ("Sources", 28), ("Queries", 9),
                ("Listings Seen", 13), ("New Roles", 10), ("Shortlisted", 11), ("Duplicates Skipped", 18)]
    for i, (name, width) in enumerate(log_cols, start=1):
        c = log.cell(row=1, column=i, value=name)
        c.fill, c.font = head_fill, head_font
        log.column_dimensions[get_column_letter(i)].width = width
    for r, e in enumerate(reversed(st.get("runs_log", [])), start=2):
        row = [e.get("run"), dt.date.fromisoformat(e["date"]), dt.date.fromisoformat(e["since"]) if e.get("since") else None,
               e.get("sources"), e.get("queries"), e.get("seen"), e.get("new"), e.get("shortlisted"), e.get("duplicates")]
        for ci, v in enumerate(row, start=1):
            c = log.cell(row=r, column=ci, value=v)
            if isinstance(v, dt.date):
                c.number_format = "yyyy-mm-dd"
    log.freeze_panes = "A2"

    guide = wb.create_sheet("How to use")
    guide.column_dimensions["A"].width = 110
    notes = [
        "Laika, The Job Retriever - Job Tracker",
        "",
        "Every role Laika has scanned is listed once on the 'Job Tracker' sheet (newest search first).",
        "Status: Pending = shown to you, waiting for your decision | Applied | Ignored | Low Match = scanned, scored below your threshold.",
        "You can change a Status here (drop-down) or type in 'My Notes'. Save and close the file - Laika picks the edits up on its next run.",
        "Tailored CV / Cover Letter link to what Laika made for that role (in 'Tailored CVs' / 'Cover Letters').",
        "Do not edit the 'Job ID' column; Laika uses it to match rows.",
        "Laika rewrites this file after every search and decision, so keep it closed while Laika is working.",
    ]
    for i, t in enumerate(notes, start=1):
        c = guide.cell(row=i, column=1, value=t)
        if i == 1:
            c.font = Font(bold=True, size=14)

    path = m / F_TRACKER
    tmp = m / ".Job Tracker.tmp.xlsx"
    try:
        wb.save(tmp)
        os.replace(tmp, path)
    except PermissionError:
        if tmp.exists():
            try:
                tmp.unlink()
            except OSError:
                pass
        st["excel_dirty"] = True
        print(f"WARNING: {F_TRACKER} is open in another program, so it was not updated. "
              "Close it - Laika will update it on the next command (or run: laika.py excel).")
        return False
    for j in db["jobs"]:
        j["excel_status"] = j["status"]
    st["excel_dirty"] = False
    return True


# ----------------------------------------------------------------------------
# Role cards & reports
# ----------------------------------------------------------------------------
def card_text(j: dict) -> str:
    score = "-" if j.get("score") is None else f"{j['score']}/100"
    lines = [f"# {j['title']} - {j['company']}", "",
             f"- **Status:** {j['status']} (since {j.get('status_date', '-')})",
             f"- **Relevancy score:** {score}",
             f"- **Location:** {j.get('location') or '-'}",
             f"- **Source:** {j.get('source') or '-'}",
             f"- **Link:** {j.get('url') or '-'}",
             *([f"- **Tailored CV:** {j['tailored_cv']} ({j.get('tailored_cv_date', '-')})"] if j.get("tailored_cv") else []),
             *([f"- **Cover letter:** {j['cover_letter']} ({j.get('cover_letter_date', '-')})"] if j.get("cover_letter") else []),
             f"- **Posted:** {j.get('posted') or '-'}",
             f"- **Found in search:** {j['search_date']} (run {j.get('run_id', '-')})",
             f"- **Laika ID:** {j['id']}", "",
             "## Why it is relevant", j.get("why") or "-", "",
             "## Possible concerns", j.get("concerns") or "-", ""]
    if j.get("reason"):
        lines += ["## Your note", j["reason"], ""]
    det = j.get("details") or {}
    meta = [f"- {k.replace('_', ' ').title()}: {v}" for k, v in det.items() if k not in ("description", "closed") and v]
    if meta:
        lines += ["## Posting details"] + meta + [""]
    if det.get("description"):
        lines += ["## Job description (snapshot taken at search time)", det["description"], ""]
    return "\n".join(lines)


def remove_cards(m: Path, j: dict) -> None:
    tag = f"[{j['id']}].md"
    for d in CARD_DIRS.values():
        folder = m / d
        if folder.exists():
            for p in folder.iterdir():
                if p.name.endswith(tag):
                    p.unlink()


def write_card(m: Path, j: dict) -> None:
    remove_cards(m, j)
    folder = CARD_DIRS.get(j["status"])
    if not folder:
        return
    target = m / folder
    target.mkdir(parents=True, exist_ok=True)
    # Stay well inside the Windows 260-character path limit.
    budget = max(0, 240 - len(str(target)) - len(f"\\{j['search_date']} -  -  [{j['id']}].md"))
    comp_len = min(30, budget // 3)
    title_len = min(50, budget - comp_len)
    if comp_len >= 4 and title_len >= 6:
        name = f"{j['search_date']} - {safe_name(j['company'], comp_len)} - {safe_name(j['title'], title_len)} [{j['id']}].md"
    else:
        name = f"{j['search_date']} [{j['id']}].md"
    try:
        (target / name).write_text(card_text(j), encoding="utf-8")
    except OSError as e:
        print(f"WARNING: could not write the role card for {j['id']} ({e}).")


def md_cell(s) -> str:
    return str(s if s is not None else "-").replace("|", "/").replace("\n", " ")


def write_report(m: Path, run: dict, jobs: list[dict], shortlist: list[dict], stats: dict, threshold: int) -> Path:
    started = dt.datetime.fromisoformat(run["started_at"])
    path = m / DIR_SCANNED / f"{started:%Y-%m-%d %H%M} - Search Report.md"
    sources = sorted({q["site"] for q in run["plan"] if q["qid"] not in run.get("skipped", [])})
    lines = [f"# Search report - {started:%Y-%m-%d %H:%M}", "",
             f"- **Window:** roles posted {window_text(run)} - each query starts from the day it last ran",
             f"- **Sources:** {', '.join(sources) or '-'}",
             f"- **Listings seen:** {stats['seen']}  |  **New roles registered:** {len(jobs)}  |  "
             f"**Duplicates skipped:** {stats['duplicates']}",
             f"- **Shortlisted (score >= {threshold}):** {len(shortlist)}", ""]
    if run.get("skipped"):
        lines += [f"- **Skipped queries:** {', '.join(run['skipped'])}", ""]
    lines += ["## Queries", "", "| ID | Site | Role | Location(s) | Since | Status |", "|---|---|---|---|---|---|"]
    for q in run["plan"]:
        loc = q.get("location") or ", ".join(q.get("locations", []))
        state = "skipped" if q["qid"] in run.get("skipped", []) else "searched"
        lines.append(f"| {q['qid']} | {md_cell(q['site'])} | {md_cell(q['role'])} | {md_cell(loc)} | {q['since']} | {state} |")
    lines += ["", "## Shortlist", ""]
    if not shortlist:
        lines.append("_No new roles above the threshold this time._")
    for n, j in enumerate(shortlist, start=1):
        lines += [f"### {n}. {j['title']} - {j['company']} ({j['score']}/100)",
                  f"- {j.get('location') or '-'} | {j.get('source')} | posted {j.get('posted') or '-'} | {j['id']}",
                  f"- {j['url']}", "", f"**Why:** {j.get('why') or '-'}", "",
                  f"**Concerns:** {j.get('concerns') or '-'}", ""]
    lines += ["## All roles scanned in this search", "",
              "| Score | Status | Role | Company | Location | Source | Link | ID |",
              "|---|---|---|---|---|---|---|---|"]
    for j in sorted(jobs, key=lambda x: -(x["score"] if x.get("score") is not None else -1)):
        lines.append(f"| {md_cell(j.get('score'))} | {j['status']} | {md_cell(j['title'])} | {md_cell(j['company'])} | "
                     f"{md_cell(j.get('location'))} | {md_cell(j.get('source'))} | [open]({j['url']}) | {j['id']} |")
    (m / DIR_SCANNED).mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return path


# ----------------------------------------------------------------------------
# Shared sync step (Excel edits in, dirty Excel out)
# ----------------------------------------------------------------------------
def sync(m: Path, db: dict, st: dict, quiet: bool = False) -> list[tuple[dict, str, str]]:
    changed = excel_sync_in(m, db)
    for j, _, _ in changed:
        write_card(m, j)
    if changed and not quiet:
        print("Imported status edits you made in Excel:")
        for j, old, new in changed:
            print(f"  {j['id']} {j['title']} - {j['company']}: {old} -> {new}")
    if changed:
        st["decisions_since_review"] = st.get("decisions_since_review", 0) + len(changed)
    if changed or st.get("excel_dirty"):
        excel_write(m, db, st)  # also records what Excel now shows, so it must be saved below
        save_db(m, db)
        save_state(m, st)
    return changed


def print_job_line(n, j: dict) -> None:
    score = "--" if j.get("score") is None else f"{j['score']:>2}"
    print(f"#{n:<3} [{j['id']}] {score}  {j['title']} | {j['company']} | {j.get('location') or '-'} | "
          f"{j.get('source')} | posted {j.get('posted') or '-'}")
    print(f"      {j['url']}")


# ----------------------------------------------------------------------------
# Commands
# ----------------------------------------------------------------------------
def cmd_init(a):
    master = Path(a.master).expanduser().resolve()
    existing = (master / F_STATE).exists()
    # Never overwrite someone's existing Laika profile by accident (e.g. onboarding pointed at an old folder).
    keep_profile = existing and not a.replace_profile
    master.mkdir(parents=True, exist_ok=True)
    for d in ALL_DIRS:
        (master / d).mkdir(parents=True, exist_ok=True)
    if not (master / F_LOCATIONS).exists() or (a.locations and not keep_profile):
        write_list(master, "locations", a.locations or [])
    if not (master / F_ROLES).exists() or (a.roles and not keep_profile):
        write_list(master, "roles", a.roles or [])
    if not (master / F_SITES).exists():
        write_list(master, "sites", DEFAULT_SITES)
    if not (master / F_CRITERIA).exists():
        write_criteria(master, dict(CRITERIA_DEFAULTS))
    if not (master / F_CHANGELOG).exists():
        (master / F_CHANGELOG).write_text("# Search criteria change log\n\n", encoding="utf-8")
    if not (master / F_DB).exists():
        save_db(master, {"version": 1, "next_id": 1, "jobs": []})
    st = load_state(master) if existing else {"created": now_local().isoformat(), "version": VERSION}
    if a.cv and not keep_profile:
        copy_cv(master, Path(a.cv))
    if keep_profile and (a.cv or a.roles or a.locations):
        db = load_db(master)
        print(f"EXISTING_SETUP: this folder already holds a Laika profile ({len(read_lines(master / F_ROLES))} roles, "
              f"{len(read_lines(master / F_LOCATIONS))} locations, {len(db['jobs'])} tracked roles). It was kept "
              "unchanged - the CV, roles and locations you passed were NOT applied. Ask the user whether to use this "
              "existing profile or choose a different folder (only with their explicit OK: --replace-profile).")
    cfg = load_config()
    cfg["master"] = str(master)
    save_config(cfg)
    if not existing:
        log_change(master, "Laika set up this master folder")
    save_state(master, st)
    print(("Re-using existing" if existing else "Created") + f" master folder: {master}")
    for d in ALL_DIRS:
        if d != DIR_INTERNAL:
            print(f"  {d}/")
    print(f"  {F_TRACKER} (created after the first search)")
    print(f"Config saved to {CONFIG_PATH}")


def copy_cv(m: Path, src: Path) -> Path:
    if not src.exists():
        die(f"CV file not found: {src}")
    cvdir = m / DIR_CV
    cvdir.mkdir(parents=True, exist_ok=True)
    old = cv_files(m)
    dest = cvdir / src.name
    if src.resolve() == dest.resolve():
        return dest
    if old:
        arch = cvdir / "Archive"
        arch.mkdir(exist_ok=True)
        for p in old:
            shutil.move(str(p), str(arch / f"{dt.date.today():%Y-%m-%d} {p.name}"))
    shutil.copy2(src, dest)
    print(f"CV saved: {dest}" + (f"  (previous version moved to {DIR_CV}/Archive)" if old else ""))
    return dest


def cmd_cv(a):
    m = get_master()
    if a.action == "add":
        copy_cv(m, Path(a.file).expanduser())
        return
    files = cv_files(m)
    if not files:
        die(f"No CV found in {m / DIR_CV}")
    cur = max(files, key=lambda p: p.stat().st_mtime)
    if a.action == "path":
        print(cur)
        return
    text = extract_text(cur)
    if text is None:
        print(f"CV_FILE: {cur}")
        print("Could not extract text automatically (install pypdf for PDFs: python -m pip install pypdf). "
              "Read the file directly instead.")
        return
    (m / DIR_INTERNAL / "cv_text.txt").write_text(text, encoding="utf-8")
    print(f"CV_FILE: {cur}\n" + "-" * 60)
    print(text)
    certs = cert_files(m)
    if certs:
        print("-" * 60 + f"\nLICENSES & CERTIFICATIONS: {len(certs)} document(s) in {DIR_CERTS}/ - read them too "
              "(`certs text`).")


def cmd_certs(a):
    m = get_master()
    d = m / DIR_CERTS
    d.mkdir(parents=True, exist_ok=True)
    if a.action == "add":
        if not a.files:
            die("certs add needs one or more file paths")
        for f in a.files:
            src = Path(f).expanduser()
            if not src.is_file():
                die(f"File not found: {src}")
            dest = d / src.name
            if src.resolve() == dest.resolve():
                print(f"Already in place: {dest}")
                continue
            existed = dest.exists()
            shutil.copy2(src, dest)
            print(f"{'Replaced' if existed else 'Saved'}: {dest}")
        return
    files = cert_files(m)
    if not files:
        print(f"No documents in {DIR_CERTS}/ - that's fine, Laika works from the CV alone. "
              "Licenses and certificates are optional (PDF, Word, text or image files).")
        return
    for p in files:
        rel = p.relative_to(d).as_posix()
        if a.action == "list":
            print(f"  {rel} ({p.suffix.lower().lstrip('.') or 'file'}, {max(1, p.stat().st_size // 1024)} KB)")
            continue
        print(f"=== {rel}")
        if p.suffix.lower() in IMAGE_EXT:
            print(f"IMAGE - no text layer. Read the file directly to see the certificate: {p}")
            continue
        try:
            text = extract_text(p)
        except Exception as e:  # noqa: BLE001 - one unreadable file should not hide the others
            text, err = None, e
        else:
            err = None
        if text is None:
            why = f"error: {err}" if err else ("install pypdf for PDFs" if p.suffix.lower() == ".pdf" else "unsupported type")
            print(f"Could not extract text ({why}). Read the file directly: {p}")
        elif not text.strip():
            print(f"No text layer (scanned document?). Read the file directly: {p}")
        else:
            print(trim(text, 4000))


def load_cv_builder():
    """The CV builder module next to this script (kept separate so it also runs on its own)."""
    sys.dont_write_bytecode = True  # keep the skill folder free of __pycache__
    here = str(Path(__file__).resolve().parent)
    if here not in sys.path:
        sys.path.insert(0, here)
    import cv_builder  # noqa: E402
    return cv_builder


def month_tag(d: dt.date | None = None) -> str:
    """The standard CV file-name date, e.g. Oct26 or Sept26."""
    d = d or dt.date.today()
    return f"{load_cv_builder().MONTH_TAGS[d.month - 1]}{d:%y}"


def role_dir(m: Path, j: dict, top: str = DIR_TAILORED) -> Path:
    """The role's folder under Tailored CVs/ or Cover Letters/ (found again by its [Job ID])."""
    base = m / top
    if base.exists():
        for p in base.iterdir():
            if p.is_dir() and p.name.endswith(f"[{j['id']}]"):
                return p
    # Leave room inside the Windows path limit for the CV file names in the folder.
    budget = max(0, 215 - len(str(base)) - len(f"\\ -  [{j['id']}]"))
    comp_len = min(30, budget // 3)
    title_len = min(50, budget - comp_len)
    if comp_len >= 4 and title_len >= 6:
        return base / f"{safe_name(j['company'], comp_len)} - {safe_name(j['title'], title_len)} [{j['id']}]"
    return base / f"[{j['id']}]"


def cmd_tailor(a):
    m = get_master()
    st, db = load_state(m), load_db(m)
    jobs = resolve_refs(a.ref, db, st)
    if len(jobs) != 1:
        die("tailor takes one role at a time: a list number or a job ID like J0012")
    j = jobs[0]
    folder = role_dir(m, j)
    folder.mkdir(parents=True, exist_ok=True)
    cvs = cv_files(m)
    cur = max(cvs, key=lambda p: p.stat().st_mtime) if cvs else None
    certs = cert_files(m)
    base = m / F_CV_BASE
    content = folder / "content.json"
    print(f"TAILOR FOLDER: {folder}")
    print(f"CONTENT FILE: {content} ({'exists - update it' if content.exists() else 'write it'})")
    if not base.exists():
        print(f"BASE CONTENT: {base} - not created yet. Write it first from the CV and certificates "
              "(see references/tailored-cv.md).")
    else:
        newer = [p.name for p in ([cur] if cur else []) + certs if p.stat().st_mtime > base.stat().st_mtime]
        print(f"BASE CONTENT: {base}" + (f" - OLDER than {', '.join(newer)}: refresh it first." if newer else ""))
    print(f"CV: {cur or 'MISSING'}")
    print(f"LICENSES & CERTIFICATIONS: {len(certs)} document(s)" +
          (": " + ", ".join(cert_rel(m, p) for p in certs) if certs else " (optional - CV only)"))
    print(f"OUTPUT NAME: First_Last_CV_{month_tag()} (cv-build derives it from the name in content.json)")
    if j.get("tailored_cv"):
        print(f"EXISTING TAILORED CV: {j['tailored_cv']} ({j.get('tailored_cv_date', '-')})")
    print(f"BUILD: laika cv-build \"{folder}\"")
    print("=" * 70)
    print(card_text(j))


def cmd_cv_build(a):
    m = get_master()
    src = Path(a.target).expanduser()
    content = src / "content.json" if src.is_dir() else src
    if not content.is_file():
        die(f"content.json not found: {content}")
    try:
        cv = json.loads(content.read_text(encoding="utf-8-sig"))
    except json.JSONDecodeError as e:
        die(f"{content} is not valid JSON: {e}")
    cv_builder = load_cv_builder()
    person = re.sub(r"[^\w-]+", "_", safe_name(cv.get("name") or "CV")).strip("_") or "CV"
    name = a.name or f"{person}_CV_{month_tag()}"
    out_base = content.parent / name
    code = cv_builder.build(content, out_base, pdf=a.pdf)
    if code == 0:
        record_document(m, content.parent, out_base, "tailored_cv")
    if code:
        sys.exit(code)


def style_state(m: Path) -> str:
    f = m / F_STYLE
    if not f.exists():
        return "NOT SET - ask the user before the first cover letter (see references/cover-letter.md)"
    text = f.read_text(encoding="utf-8", errors="replace")
    hit = re.search(r"\*\*Mode:\*\*\s*([^.\n]+)", text)
    samples = len(style_samples(m))
    return (hit.group(1).strip() if hit else "set") + (f" | {samples} sample(s)" if samples else "")


def style_fit(m: Path) -> str | None:
    """The style file's application-fit decision, or None when it hasn't been checked yet."""
    f = m / F_STYLE
    if not f.exists():
        return None
    text = f.read_text(encoding="utf-8", errors="replace")
    hit = re.search(r"\*\*Application fit:\*\*\s*([^\n_]+)", text)
    if hit:
        return hit.group(1).strip().rstrip(".")
    if re.search(r"\*\*Mode:\*\*\s*default professional", text):
        return "professional by default"
    return None


def style_samples(m: Path) -> list[Path]:
    d = m / DIR_STYLE_SAMPLES
    if not d.exists():
        return []
    return sorted(p for p in d.rglob("*")
                  if p.is_file() and not any(part.startswith(("~$", ".")) for part in p.relative_to(d).parts))


def eml_text(p: Path) -> str:
    import email
    from email import policy
    msg = email.message_from_bytes(p.read_bytes(), policy=policy.default)
    part = msg.get_body(preferencelist=("plain", "html"))
    body = ""
    if part is not None:
        body = part.get_content()
        if part.get_content_type() == "text/html":
            body = html_to_text(body)
    return f"From: {msg['from']}\nTo: {msg['to']}\nSubject: {msg['subject']}\n\n{body}"


def cmd_style(a):
    m = get_master()
    d = m / DIR_STYLE_SAMPLES
    if a.action == "add":
        if not a.files:
            die("style add needs one or more file paths")
        d.mkdir(parents=True, exist_ok=True)
        for f in a.files:
            src = Path(f).expanduser()
            if not src.is_file():
                die(f"File not found: {src}")
            dest = d / src.name
            existed = dest.exists()
            if src.resolve() != dest.resolve():
                shutil.copy2(src, dest)
            print(f"{'Replaced' if existed else 'Saved'}: {dest}")
        return
    if a.action == "path":
        print(m / F_STYLE)
        return
    if a.action == "show":
        print(f"WRITING STYLE: {style_state(m)}")
        print(f"FILE: {m / F_STYLE}")
        if (m / F_STYLE).exists():
            print("-" * 60 + "\n" + (m / F_STYLE).read_text(encoding="utf-8"))
        return
    files = style_samples(m)  # action == "samples"
    if not files:
        print(f"No writing samples in {DIR_STYLE_SAMPLES}/.")
        return
    for p in files:
        print(f"=== {p.relative_to(d).as_posix()}")
        if p.suffix.lower() in IMAGE_EXT:
            print(f"IMAGE - Read the file directly: {p}")
            continue
        try:
            text = eml_text(p) if p.suffix.lower() == ".eml" else extract_text(p)
        except Exception as e:  # noqa: BLE001 - one unreadable file should not hide the others
            text = None
            print(f"(error: {e})")
        if text is None:
            print(f"Could not extract text. Read the file directly: {p}")
        else:
            print(trim(text, a.chars))


def cmd_letter(a):
    m = get_master()
    st, db = load_state(m), load_db(m)
    jobs = resolve_refs(a.ref, db, st)
    if len(jobs) != 1:
        die("letter takes one role at a time: a list number or a job ID like J0012")
    j = jobs[0]
    folder = role_dir(m, j, DIR_LETTERS)
    folder.mkdir(parents=True, exist_ok=True)
    content = folder / "letter.json"
    base = m / F_CV_BASE
    certs = cert_files(m)
    print(f"LETTER FOLDER: {folder}")
    print(f"LETTER FILE: {content} ({'exists - revise it' if content.exists() else 'write it'})")
    print(f"WRITING STYLE: {style_state(m)} -> {m / F_STYLE}")
    if (m / F_STYLE).exists():
        print(f"APPLICATION FIT: {style_fit(m) or 'NOT CHECKED - check the style against job applications first and ask the user if it is too casual (references/cover-letter.md, Application fit)'}")
    print(f"FACTS: {base if base.exists() else 'Base CV content not created yet - write it first (references/tailored-cv.md)'}")
    print(f"LICENSES & CERTIFICATIONS: {len(certs)} document(s)" + ("" if certs else " (optional - CV only)"))
    if j.get("tailored_cv"):
        print(f"TAILORED CV FOR THIS ROLE: {j['tailored_cv']} (keep the letter consistent with it)")
    if j.get("cover_letter"):
        print(f"EXISTING COVER LETTER: {j['cover_letter']} ({j.get('cover_letter_date', '-')})")
    print(f"OUTPUT NAME: First_Last_Cover_Letter_{month_tag()} (letter-build derives it from the name in letter.json)")
    print(f'BUILD: laika letter-build "{folder}"')
    print("=" * 70)
    print(card_text(j))


def record_document(m: Path, folder: Path, out_base: Path, key: str) -> None:
    """Link a built CV or letter on its role (folder name ends in [Job ID]): db, card and Excel."""
    hit = re.search(r"\[(J\d+)\]$", folder.name)
    if not hit:
        return
    db, st = load_db(m), load_state(m)
    j = find_job(db, hit.group(1))
    if not j:
        return
    pdf, docx = Path(f"{out_base}.pdf"), Path(f"{out_base}.docx")
    j[key] = (pdf if pdf.exists() else docx).resolve().relative_to(m.resolve()).as_posix()
    j[f"{key}_date"] = today_str()
    write_card(m, j)
    excel_ok = excel_write(m, db, st)
    save_db(m, db)
    save_state(m, st)
    print(f"Recorded on {j['id']} ({j['title']} - {j['company']}); Excel {'updated' if excel_ok else 'NOT updated (file open)'}.")


def cmd_letter_build(a):
    m = get_master()
    src = Path(a.target).expanduser()
    content = src / "letter.json" if src.is_dir() else src
    if not content.is_file():
        die(f"letter.json not found: {content}")
    try:
        letter = json.loads(content.read_text(encoding="utf-8-sig"))
    except json.JSONDecodeError as e:
        die(f"{content} is not valid JSON: {e}")
    cv_builder = load_cv_builder()
    person = re.sub(r"[^\w-]+", "_", safe_name(letter.get("name") or "Cover")).strip("_") or "Cover"
    out_base = content.parent / (a.name or f"{person}_Cover_Letter_{month_tag()}")
    code = cv_builder.build_letter(content, out_base, pdf=a.pdf)
    if code == 0:
        record_document(m, content.parent, out_base, "cover_letter")
    if code:
        sys.exit(code)


def cmd_status(a):
    m = get_master(required=False)
    if not m:
        cfg = load_config()
        print("STATE: NOT_INITIALIZED")
        print(f"Profile: {CONFIG_KIND} ({CONFIG_PATH})")
        if cfg.get("master"):
            print(f"Configured master folder is missing or empty: {cfg['master']} (moved or deleted?)")
        print("Start onboarding: ask for CV, locations, roles, then the master folder.")
        return
    st, db = load_state(m), load_db(m)
    sync(m, db, st)
    crit = read_criteria(m)
    print("STATE: READY")
    print(f"Profile: {CONFIG_KIND} ({CONFIG_PATH})")
    print(f"Master folder: {m}")
    cvs = cv_files(m)
    print(f"CV: {', '.join(p.name for p in cvs) if cvs else 'MISSING'}"
          f" | Laika CV summary: {'yes' if (m / F_CV_SUMMARY).exists() else 'NOT WRITTEN YET'}")
    for d in (DIR_CERTS, DIR_TAILORED, DIR_LETTERS):  # folders added after 1.0; older setups get them here
        (m / d).mkdir(exist_ok=True)
    certs = cert_files(m)
    names = ", ".join(p.name for p in certs[:6]) + (f" (+{len(certs) - 6} more)" if len(certs) > 6 else "")
    print(f"Licenses & certifications: {f'{len(certs)} document(s): {names}' if certs else 'none (optional)'}")
    print(f"Writing style (cover letters): {style_state(m)}")
    roles, locs, sites = read_lines(m / F_ROLES), read_lines(m / F_LOCATIONS), read_sites(m)
    print(f"Roles ({len(roles)}): {'; '.join(roles) or '-'}")
    print(f"Locations ({len(locs)}): {'; '.join(locs) or '-'}")
    site_text = "; ".join(f"{s['name']} [{s['method']}]" for s in sites) or "-"
    print(f"Job sites ({len(sites)}): {site_text}")
    print("Criteria: " + " | ".join(f"{k}={v}" for k, v in crit.items() if v not in ("", None)))
    changes = detect_changes(m, st)
    if not st.get("snapshot"):
        print("PROFILE CHANGES: no baseline yet (first run) - run `snapshot` once the profile is confirmed.")
    elif changes:
        print("PROFILE CHANGES since the last run:")
        for c in changes:
            print(f"  * {c['file']}: {c['change']}")
            for ln in c.get("added_lines", []):
                print(f"      + {ln}")
            for ln in c.get("removed_lines", []):
                print(f"      - {ln}")
    else:
        print("PROFILE CHANGES: none")
    if st.get("last_search_at"):
        print(f"Last search: {st['last_search_at'].replace('T', ' ')} (run {st.get('last_search_run')})")
    else:
        print(f"Last search: never - first search will look back {crit_int(crit, 'first_search_lookback_days')} days")
    counts = {s: 0 for s in STATUSES}
    for j in db["jobs"]:
        counts[j["status"]] = counts.get(j["status"], 0) + 1
    print(f"Tracker: {len(db['jobs'])} roles | " + " | ".join(f"{s} {n}" for s, n in counts.items()))
    print(f"Decisions since last criteria review: {st.get('decisions_since_review', 0)}")
    run = load_run(m, st, required=False)
    if run and not run.get("registered"):
        print(f"OPEN RUN: {run['id']} started {run['started_at'].replace('T', ' ')} with "
              f"{len(run['candidates'])} candidates (not registered yet)")
    if st.get("excel_dirty"):
        print("Excel: NOT up to date (file was open) - close it and run `excel`")
    if a.verbose and (m / F_CV_SUMMARY).exists():
        print("-" * 60 + "\nLAIKA CV SUMMARY:\n" + (m / F_CV_SUMMARY).read_text(encoding="utf-8"))


def cmd_snapshot(a):
    m = get_master()
    st = load_state(m)
    take_snapshot(m, st)
    save_state(m, st)
    print(f"Profile baseline saved ({len(st['snapshot'])} files). Future runs report changes against it.")


def cmd_list(a):
    m = get_master()
    st = load_state(m)
    rel = LIST_SPECS[a.kind][0]
    items = read_lines(m / rel)
    if a.action == "show":
        print(f"{a.kind} ({len(items)}):")
        for it in items:
            print(f"  - {it}")
        return
    before = list(items)
    if a.action == "set":
        items = a.items
    elif a.action == "add":
        items = items + a.items
    elif a.action == "remove":
        drop = {norm(x) for x in a.items}
        items = [x for x in items if norm(x) not in drop]
        missing = drop - {norm(x) for x in before}
        if missing:
            print(f"Note: not found in {a.kind}: {', '.join(sorted(missing))}")
    write_list(m, a.kind, items)
    after = read_lines(m / rel)
    added = [x for x in after if norm(x) not in {norm(b) for b in before}]
    removed = [x for x in before if norm(x) not in {norm(b) for b in after}]
    if added or removed:
        desc = f"{a.kind}: " + "; ".join(([f"added {added}"] if added else []) + ([f"removed {removed}"] if removed else []))
        log_change(m, desc, a.reason)
        print("Updated " + desc)
    else:
        print(f"No change to {a.kind}.")
    take_snapshot(m, st, only=[rel])
    save_state(m, st)


def cmd_criteria(a):
    m = get_master()
    st = load_state(m)
    crit = read_criteria(m)
    if a.action == "show":
        for k, _, desc in CRITERIA_SPEC:
            print(f"{k}: {crit.get(k, '')}    # {desc}")
        for k in crit:
            if k not in CRITERIA_DEFAULTS:
                print(f"{k}: {crit[k]}    # custom")
        return
    key = a.key.strip().lower().replace(" ", "_")
    if key in ("min_relevancy_score", "max_results_per_query", "first_search_lookback_days"):
        if not a.value.strip().isdigit():
            die(f"{key} must be a whole number")
    old = crit.get(key, "")
    crit[key] = a.value.strip()
    write_criteria(m, crit)
    if old != crit[key]:
        log_change(m, f"criteria '{key}': '{old}' -> '{crit[key]}'", a.reason)
        print(f"Updated {key}: '{old}' -> '{crit[key]}'")
    else:
        print(f"{key} unchanged.")
    take_snapshot(m, st, only=[F_CRITERIA])
    save_state(m, st)


def cmd_run(a):
    m = get_master()
    st = load_state(m)
    if a.action == "start":
        cur = load_run(m, st, required=False)
        if cur and not cur.get("registered") and not a.new:
            print(f"RESUMING open run {cur['id']} (use `run start --new` to discard it).")
            print_plan(cur)
            return
        crit = read_criteria(m)
        if not read_lines(m / F_ROLES) or not read_lines(m / F_LOCATIONS) or not read_sites(m):
            die("roles, locations and job sites must each have at least one entry before searching.")
        now = now_local()
        rid = now.strftime("%Y-%m-%d_%H%M")
        k = 2
        while run_path(m, rid).exists():
            rid = f"{now:%Y-%m-%d_%H%M}-{k}"
            k += 1
        run = {"id": rid, "started_at": now.isoformat(), "plan": compute_plan(m, st, crit, now),
               "candidates": [], "stats": {"seen": 0, "duplicates": 0, "old": 0}, "registered": False}
        save_run(m, run)
        st["current_run"] = rid
        save_state(m, st)
        print(f"RUN STARTED: {rid}")
        print_plan(run)
        return
    run = load_run(m, st)
    if a.action == "status":
        print_plan(run)
        by_src: dict = {}
        for c in run["candidates"]:
            by_src[c["source"]] = by_src.get(c["source"], 0) + 1
        print(f"Candidates: {len(run['candidates'])} " + str(by_src) +
              f" | seen {run['stats']['seen']} | duplicates skipped {run['stats']['duplicates']}")
    elif a.action == "done":
        ids = {x.strip() for x in ",".join(a.qids).split(",") if x.strip()}
        for q in run["plan"]:
            if q["qid"] in ids:
                q["done"] = True
        save_run(m, run)
        print(f"Marked done: {', '.join(sorted(ids))}")
    elif a.action == "abandon":
        run["abandoned"] = True
        run["abandoned_at"] = now_local().isoformat()
        save_run(m, run)
        st["current_run"] = None
        save_state(m, st)
        print(f"Run {run['id']} discarded. Nothing was registered, and the run can't be resumed or re-linked.")


def print_plan(run: dict) -> None:
    print(f"Search plan for run {run['id']} ({len(run['plan'])} queries):")
    for method, label in (("api", "LinkedIn - fetched by `fetch-linkedin`"),
                          ("browser", "Browser (built-in browser or Claude in Chrome) - search in-page, keep matching locations"),
                          ("web", "Web search - use the search hint, then fetch the posting pages")):
        qs = [q for q in run["plan"] if q["method"] == method]
        if not qs:
            continue
        print(f"\n[{label}]")
        for q in qs:
            loc = q.get("location") or " / ".join(q.get("locations", []))
            flag = " (new combination - full lookback)" if q.get("new_combo") else ""
            done = " [done]" if q.get("done") else (f" [stopped early: {q['error']}]" if q.get("error") else "")
            print(f"  {q['qid']} {q['site']}: {q['role']} @ {loc} | since {q['since']}{flag}{done}")
            if method != "api":
                print(f"      {q.get('search_hint') or q['url']}")


def cmd_fetch_linkedin(a):
    m = get_master()
    st, db, crit = load_state(m), load_db(m), read_criteria(m)
    run = load_run(m, st)
    roles = match_terms(m)
    idx = db_index(db)
    maxq = a.max_per_query or crit_int(crit, "max_results_per_query")
    only = {x.strip() for x in (a.queries or "").split(",") if x.strip()}
    qs = [q for q in run["plan"] if q["method"] == "api" and (not only or q["qid"] in only)]
    if not qs:
        print("No LinkedIn queries in this run's plan.")
        return
    todo = qs if a.refetch else [q for q in qs if not q.get("done")]
    if len(todo) < len(qs):
        print(f"{len(qs) - len(todo)} queries already fetched - skipping them (use --refetch to repeat).")
    total_new = 0
    for q in todo:
        seen = new = dup = old = 0
        err = None
        for start in range(0, maxq, 10):
            url = ("https://www.linkedin.com/jobs-guest/jobs/api/seeMoreJobPostings/search?" +
                   urllib.parse.urlencode(linkedin_params(q, start)))
            try:
                page = http_get(url)
            except FetchError as e:
                err = str(e)
                break
            cards = parse_linkedin_cards(page)
            if not cards:
                break
            for card in cards[: maxq - seen]:
                seen += 1
                outcome, _ = add_candidate(run, card, idx, crit, roles, "LinkedIn",
                                           f"{q['role']} @ {q['location']}", q["since"])
                if outcome == "new":
                    new += 1
                elif outcome in ("dup_db", "dup_run", "merged"):
                    dup += 1
                elif outcome == "old":
                    old += 1
            if len(cards) < 10 or seen >= maxq:
                break
            polite_sleep(a.delay)
        run["stats"]["seen"] += seen
        run["stats"]["duplicates"] += dup
        run["stats"]["old"] = run["stats"].get("old", 0) + old
        if err:
            # Partial results stay in the run, but the query is not done (also after --refetch): a plain rerun
            # fetches it again, and if it never finishes, `register` keeps its old window.
            q["done"] = False
            q["error"] = err
            print(f"{q['qid']} {q['role']} @ {q['location']}: stopped early ({err}) - seen {seen}, new {new} "
                  "(kept, not marked done)")
        else:
            q["done"] = True
            q.pop("error", None)
            print(f"{q['qid']} {q['role']} @ {q['location']}: seen {seen}, new {new}, duplicates {dup}, "
                  f"older than window {old}")
        total_new += new
        save_run(m, run)
        polite_sleep(a.delay)
    incomplete = [f"{q['qid']} ({q.get('error') or 'not fetched'})" for q in qs if not q.get("done")]
    if incomplete:
        print("Some queries stopped early (often LinkedIn rate limiting). What they found is kept, but they are "
              "not marked done. Wait a few minutes and run `fetch-linkedin` again - only unfinished queries run.")
        print(f"LinkedIn incomplete: {total_new} new candidates. Run total: {len(run['candidates'])}. "
              f"Not done: {', '.join(incomplete)}.")
    else:
        print(f"LinkedIn done: {total_new} new candidates. Run total: {len(run['candidates'])}.")


def cmd_add_candidates(a):
    m = get_master()
    st, db, crit = load_state(m), load_db(m), read_criteria(m)
    run = load_run(m, st)
    data = read_json(Path(a.input), None)
    if data is None:
        die(f"Could not read {a.input}")
    items = data if isinstance(data, list) else data.get("jobs", [])
    qids = {x.strip() for x in (a.query or "").split(",") if x.strip()}
    plan_q = [q for q in run["plan"] if q["qid"] in qids]
    min_since = min((q["since"] for q in plan_q), default=None) or run_min_since(run, a.source)
    idx = db_index(db)
    roles = match_terms(m)
    counts: dict = {}
    for it in items:
        src = it.get("source") or a.source
        outcome, detail = add_candidate(run, it, idx, crit, roles, src, it.get("query", ""), None if a.no_date_filter else min_since)
        counts[outcome] = counts.get(outcome, 0) + 1
        if outcome != "new":
            print(f"  skipped ({outcome}: {detail}): {it.get('title')} - {it.get('company')}")
    run["stats"]["seen"] += len(items)
    run["stats"]["duplicates"] += counts.get("dup_db", 0) + counts.get("dup_run", 0) + counts.get("merged", 0)
    run["stats"]["old"] = run["stats"].get("old", 0) + counts.get("old", 0)
    for q in plan_q:
        q["done"] = True
    save_run(m, run)
    print(f"Added {counts.get('new', 0)} new candidates from {a.source} "
          f"(duplicates {counts.get('dup_db', 0) + counts.get('dup_run', 0)}, merged {counts.get('merged', 0)}, "
          f"older than window {counts.get('old', 0)}, invalid {counts.get('invalid', 0)}). "
          f"Run total: {len(run['candidates'])}.")


def cmd_details(a):
    m = get_master()
    st = load_state(m)
    run = load_run(m, st)
    wanted = {x.strip() for x in (a.cids or "").split(",") if x.strip()}
    pool = [c for c in run["candidates"] if c["key"].startswith("linkedin:") and
            not (c.get("details") or {}).get("description") and
            (c["cid"] in wanted if wanted else (not c["flags"] and c["title_match"] >= a.min_title_match))]
    pool.sort(key=lambda c: -c["title_match"])
    pool = pool[: a.limit]
    if not pool:
        print("Nothing to fetch - all eligible LinkedIn candidates already have descriptions.")
        return
    ok = fail = 0
    for i, c in enumerate(pool, start=1):
        try:
            page = http_get("https://www.linkedin.com/jobs-guest/jobs/api/jobPosting/" + c["key"].split(":", 1)[1])
            c["details"] = parse_linkedin_detail(page) if page else {"error": "posting not available"}
            ok += 1
        except FetchError as e:
            c["details"] = {"error": str(e)}
            fail += 1
        if i % 5 == 0:
            save_run(m, run)
        polite_sleep(a.delay)
    save_run(m, run)
    skipped = len([c for c in run["candidates"] if c["key"].startswith("linkedin:") and not c.get("details")])
    print(f"Fetched descriptions for {ok} LinkedIn roles ({fail} failed). "
          f"{skipped} LinkedIn candidates still without a description (low title match or flagged).")


def cmd_candidates(a):
    m = get_master()
    st = load_state(m)
    run = load_run(m, st)
    cands = run["candidates"]
    rdir = run_path(m, run["id"]).parent
    if not a.export:
        if a.compact:  # triage view for big runs: one short line per candidate, grouped by location
            rows = [f"{c['cid']}|{c['title'][:62]}|{c['company'][:24]}|{c['location'][:26]}|{c['source'][:10]}"
                    for c in sorted(cands, key=lambda c: (norm_city(c["location"]), -c["title_match"], c["title"]))]
        else:
            rows = []
            for c in sorted(cands, key=lambda c: (bool(c["flags"]), -c["title_match"])):
                desc = "desc" if (c.get("details") or {}).get("description") else "no-desc"
                flag = f" FLAGS: {'; '.join(c['flags'])}" if c["flags"] else ""
                closed = " CLOSED" if (c.get("details") or {}).get("closed") else ""
                rows.append(f"  {c['cid']} tm={c['title_match']:.2f} {desc}{closed} | {c['title']} | {c['company']} | "
                            f"{c['location']} | {c['source']} | {c['posted'] or '-'}{flag}")
        header = f"Run {run['id']}: {len(cands)} candidates"
        if a.out:
            Path(a.out).write_text(header + "\n" + "\n".join(rows) + "\n", encoding="utf-8")
            print(f"{header} - list written to {a.out}")
        else:
            print(header)
            print("\n".join(rows))
        return
    out = rdir / "for_scoring.md"
    lines = [f"# Candidates to score - run {run['id']}", "",
             "Score every candidate below (0-100) against the CV and criteria. Write the scores file:",
             f"  {rdir / 'scores.json'}",
             'Format: [{"cid": "c001", "score": 82, "why": "...", "concerns": "..."}, ...]',
             'Repeat postings of the same job: [{"cid": "c014", "duplicate_of": "c009"}] - they are not registered.',
             "Candidates you leave out are registered as Low Match without a score.", ""]
    wanted = [x.strip() for x in (a.cids or "").split(",") if x.strip()]
    if wanted:
        by_cid = {c["cid"]: c for c in cands}
        missing = [x for x in wanted if x not in by_cid]
        if missing:
            die(f"Unknown candidate ids: {', '.join(missing)}")
        chosen = [by_cid[x] for x in wanted]
    else:
        chosen = [c for c in cands if a.all or (not c["flags"] and c["title_match"] > 0) or
                  (c.get("details") or {}).get("description")]
    for c in chosen:
        det = c.get("details") or {}
        desc = det.get("description", "")
        lines += [f"## {c['cid']} - {c['title']} | {c['company']}",
                  f"- Location: {c['location'] or '-'} | Work mode: {c.get('work_mode') or det.get('workplace_type') or '-'} | "
                  f"Source: {c['source']} | Posted: {c['posted'] or det.get('posted_ago') or '-'}",
                  f"- URL: {c['url']}"]
        meta = [f"{k.replace('_', ' ')}: {v}" for k, v in det.items() if k not in ("description", "error") and v]
        if meta:
            lines.append("- " + " | ".join(meta))
        if c["flags"]:
            lines.append(f"- AUTO-FLAGS: {'; '.join(c['flags'])}")
        if c.get("notes"):
            lines.append(f"- Notes: {c['notes']}")
        if not desc:
            lines += ["", "(no description fetched - judge from title/company)", ""]
        elif a.excerpt:
            lines += ["", description_excerpt(desc), ""]
        else:
            lines += ["", trim(desc, a.chars), ""]
    others = [] if wanted else [c for c in cands if c not in chosen]
    if others:
        lines += ["# Other candidates (title only)", "",
                  "Flagged, or the title is not obviously related to the target roles. Skim them: a local-language",
                  "title (e.g. 'Vertriebsleiter' = Sales Director) can still match - score it, or fetch its",
                  "description first with `details --cids <cid>`. Unscored ones are registered as Low Match.", ""]
        for c in others:
            flag = f" | FLAGS: {'; '.join(c['flags'])}" if c["flags"] else ""
            lines.append(f"- {c['cid']} | {c['title']} | {c['company']} | {c['location'] or '-'} | {c['source']}{flag}")
    out.write_text("\n".join(lines) + "\n", encoding="utf-8")
    print(f"\nEXPORTED {len(chosen)} candidates for scoring -> {out}")
    if others:
        print(f"{len(others)} more are listed title-only at the end of the file (flagged, or title unrelated to "
              "your roles). Unscored ones are registered as Low Match.")
    print(f"Write scores to: {rdir / 'scores.json'}")


def cmd_register(a):
    m = get_master()
    st, db, crit = load_state(m), load_db(m), read_criteria(m)
    run = load_run(m, st)
    if run.get("registered"):
        die("This run is already registered. Start a new one with `run start`.")
    scores_path = Path(a.scores) if a.scores else run_path(m, run["id"]).parent / "scores.json"
    raw = read_json(scores_path, None)
    if raw is None:
        die(f"Scores file not found: {scores_path}")
    if isinstance(raw, dict):
        raw = [dict(v, cid=k) for k, v in raw.items()]
    scores = {}
    for s in raw:
        cid = str(s.get("cid", "")).strip()
        if cid:
            scores[cid] = s
    unknown = [cid for cid in scores if cid not in {c["cid"] for c in run["candidates"]}]
    if unknown:
        die(f"Unknown candidate ids in scores file: {', '.join(unknown)}")
    threshold = crit_int(crit, "min_relevancy_score")
    sync(m, db, st, quiet=False)
    keys, fps = db_index(db)
    started = dt.datetime.fromisoformat(run["started_at"])
    search_date = started.date().isoformat()
    home_cities = set()
    for loc in read_lines(m / F_LOCATIONS):
        where, remote = parse_location(loc)
        home_cities.add(norm_city(where))
        if remote:
            home_cities.add("remote")
    jobs, late_dups, marked_dups = [], 0, 0
    for c in run["candidates"]:
        if c["key"] in keys or c["fingerprint"] in fps:
            late_dups += 1
            continue
        s = scores.get(c["cid"])
        if s and s.get("duplicate_of"):  # same job listed again (other country / aggregator repost)
            marked_dups += 1
            continue
        score = None
        if s and s.get("score") is not None:
            try:
                score = max(0, min(100, int(round(float(s["score"])))))
            except (TypeError, ValueError):
                die(f"Score for {c['cid']} is not a number: {s.get('score')}")
        if score is not None:
            why, concerns = (s.get("why") or "").strip(), (s.get("concerns") or "").strip()
        elif c["flags"]:
            why, concerns = "", "Auto-filtered: " + "; ".join(c["flags"])
        elif norm_city(c["location"]) in home_cities:
            why, concerns = "", ("Not scored - in one of your locations, but the title/function, seniority or industry "
                                 "is outside your target roles and rules.")
        else:
            why, concerns = "", (f"Not scored - located in {c['location'] or 'an unspecified place'}, outside your "
                                 "locations and not advertised as remote for you.")
        status = "Pending" if score is not None and score >= threshold else "Low Match"
        job = {"id": new_job_id(db), "key": c["key"], "fingerprint": c["fingerprint"],
               "title": c["title"], "company": c["company"], "location": c["location"], "url": c["url"],
               "source": c["source"], "posted": c.get("posted", ""), "work_mode": c.get("work_mode", ""),
               "search_date": search_date, "run_id": run["id"], "query": c.get("query", ""),
               "score": score, "why": why, "concerns": concerns, "status": status,
               "status_date": search_date, "reason": "", "user_notes": "",
               "details": c.get("details") or {}, "excel_status": None}
        db["jobs"].append(job)
        keys[job["key"]] = job
        fps[job["fingerprint"]] = job
        jobs.append(job)
    skipped = {x.strip() for x in (a.skipped or "").split(",") if x.strip()}
    # Only queries that actually ran move their window forward; the rest keep it, so nothing is missed.
    not_run = sorted(q["qid"] for q in run["plan"] if not q.get("done") and q["qid"] not in skipped)
    run["skipped"] = sorted(skipped | set(not_run), key=lambda x: int(x[1:]) if x[1:].isdigit() else 0)
    searched = st.setdefault("searched", {})
    for q in run["plan"]:
        if q["qid"] in skipped or not q.get("done"):
            continue
        for loc in ([q["location"]] if "location" in q else q.get("locations", [])):
            searched[combo_key(q["site"], q["role"], loc)] = run["started_at"]
    shortlist = sorted([j for j in jobs if j["status"] == "Pending"], key=lambda j: -j["score"])
    stats = {"seen": run["stats"]["seen"], "duplicates": run["stats"]["duplicates"] + late_dups + marked_dups}
    report = write_report(m, run, jobs, shortlist, stats, threshold)
    st["last_search_at"] = run["started_at"]
    st["last_search_run"] = run["id"]
    st["current_run"] = None
    st["last_list"] = {"kind": "shortlist", "ids": [j["id"] for j in shortlist], "at": now_local().isoformat()}
    st.setdefault("runs_log", []).append({
        "run": run["id"], "date": search_date, "since": run_min_since(run),
        "sources": ", ".join(sorted({q["site"] for q in run["plan"] if q["qid"] not in skipped})),
        "queries": len([q for q in run["plan"] if q["qid"] not in skipped]),
        "seen": stats["seen"], "new": len(jobs), "shortlisted": len(shortlist), "duplicates": stats["duplicates"]})
    run["registered"] = True
    run["registered_at"] = now_local().isoformat()
    excel_ok = excel_write(m, db, st)
    save_db(m, db)
    save_state(m, st)
    save_run(m, run)
    print(f"SEARCH REGISTERED - run {run['id']}")
    print(f"Window: {window_text(run)} | listings seen {stats['seen']} | new roles {len(jobs)} | "
          f"duplicates skipped {stats['duplicates']} | shortlisted (score >= {threshold}) {len(shortlist)} | "
          f"low match {len(jobs) - len(shortlist)}")
    print(f"Excel: {'updated' if excel_ok else 'NOT updated (file open)'} -> {m / F_TRACKER}")
    print(f"Report: {report}")
    stopped = {q["qid"]: q["error"] for q in run["plan"] if q["qid"] in not_run and q.get("error")}
    if stopped:
        print(f"NOTE: {len(stopped)} queries stopped early ({', '.join(f'{k} {v}' for k, v in stopped.items())}). "
              "What they found is registered, and they keep their old window, so the next search covers them again.")
    never = [x for x in not_run if x not in stopped]
    if never:
        print(f"NOTE: {len(never)} queries were never marked as searched ({', '.join(never)}). They keep their "
              "old window for next time. If they did run, mark them with `add-candidates --query` or `run done` "
              "before registering.")
    if not shortlist:
        print("\nNo new roles above the threshold this time.")
        return
    print("\nSHORTLIST (the # numbers work with `decide`):")
    for n, j in enumerate(shortlist, start=1):
        print_job_line(n, j)
        print(f"      WHY: {j['why']}")
        print(f"      CONCERNS: {j['concerns'] or '-'}")


def resolve_refs(text: str, db: dict, st: dict) -> list[dict]:
    out = []
    last_ids = (st.get("last_list") or {}).get("ids", [])
    for tok in re.split(r"[,\s]+", text or ""):
        tok = tok.strip().lstrip("#")
        if not tok:
            continue
        if re.fullmatch(r"\d+-\d+", tok):
            lo, hi = map(int, tok.split("-"))
            nums = list(range(lo, hi + 1))
        elif tok.isdigit():
            nums = [int(tok)]
        else:
            j = find_job(db, tok)
            if not j:
                die(f"Unknown job reference: {tok}")
            out.append(j)
            continue
        for n in nums:
            if not 1 <= n <= len(last_ids):
                die(f"#{n} is not in the last list shown ({len(last_ids)} items). Use a job ID like J0012.")
            out.append(find_job(db, last_ids[n - 1]))
    return out


def cmd_decide(a):
    m = get_master()
    st, db = load_state(m), load_db(m)
    sync(m, db, st)
    reasons = {}
    for r in a.reason or []:
        if "=" not in r:
            die(f"--reason must look like REF=text (got: {r})")
        ref, text = r.split("=", 1)
        for j in resolve_refs(ref, db, st):
            reasons[j["id"]] = text.strip()
    changes = []
    for status, refs in (("Applied", a.apply), ("Ignored", a.ignore), ("Pending", a.pending)):
        for j in resolve_refs(refs, db, st):
            old = j["status"]
            j["status"] = status
            j["status_date"] = today_str()
            j["status_via"] = "laika"
            if j["id"] in reasons:
                j["reason"] = reasons.pop(j["id"])
            changes.append((j, old, status))
    for jid, text in reasons.items():  # reason given without a status change
        j = find_job(db, jid)
        j["reason"] = text
        changes.append((j, j["status"], j["status"]))
    if not changes:
        die("Nothing to do - pass --apply / --ignore / --pending with list numbers or job IDs.")
    for j, _, _ in changes:
        write_card(m, j)
    st["decisions_since_review"] = st.get("decisions_since_review", 0) + len(changes)
    excel_ok = excel_write(m, db, st)
    save_db(m, db)
    save_state(m, st)
    for j, old, new in changes:
        arrow = f"{old} -> {new}" if old != new else f"{new} (note updated)"
        print(f"{j['id']} {j['title']} - {j['company']}: {arrow}" + (f"  [reason: {j['reason']}]" if j.get("reason") else ""))
    print(f"Excel: {'updated' if excel_ok else 'NOT updated (file open)'}. Role cards written to "
          f"'{DIR_APPLIED}' / '{DIR_IGNORED}'.")
    left = sum(1 for j in db["jobs"] if j["status"] == "Pending")
    print(f"Still pending: {left}")


def cmd_pending(a):
    m = get_master()
    st, db = load_state(m), load_db(m)
    sync(m, db, st)
    jobs = sorted([j for j in db["jobs"] if j["status"] == "Pending"],
                  key=lambda j: (-(j["score"] or 0), j["search_date"]))
    if not jobs:
        print("No roles are waiting for a decision.")
        return
    st["last_list"] = {"kind": "pending", "ids": [j["id"] for j in jobs], "at": now_local().isoformat()}
    save_state(m, st)
    print(f"PENDING ROLES ({len(jobs)}) - the # numbers work with `decide`:")
    for n, j in enumerate(jobs, start=1):
        print_job_line(n, j)
        print(f"      found {j['search_date']} | WHY: {trim(j['why'], 220)}")
        if a.verbose:
            print(f"      CONCERNS: {j['concerns'] or '-'}")


def cmd_show(a):
    m = get_master()
    st, db = load_state(m), load_db(m)
    for j in resolve_refs(a.ref, db, st):
        print(card_text(j))
        print("=" * 70)


def cmd_feedback(a):
    m = get_master()
    st, db, crit = load_state(m), load_db(m), read_criteria(m)
    sync(m, db, st)
    applied = [j for j in db["jobs"] if j["status"] == "Applied"]
    ignored = [j for j in db["jobs"] if j["status"] == "Ignored"]
    pending = [j for j in db["jobs"] if j["status"] == "Pending"]
    threshold = crit_int(crit, "min_relevancy_score")

    def avg(js):
        sc = [j["score"] for j in js if j.get("score") is not None]
        return f"{sum(sc) / len(sc):.0f}" if sc else "-"

    print("FEEDBACK SUMMARY")
    print(f"Applied {len(applied)} (avg score {avg(applied)}) | Ignored {len(ignored)} (avg score {avg(ignored)}) | "
          f"Pending {len(pending)} | threshold {threshold}")
    print(f"Decisions since last criteria review: {st.get('decisions_since_review', 0)}")

    def tok_counts(js):
        cnt: dict = {}
        for j in js:
            for t in set(tokens(j["title"])):
                if len(t) > 2:
                    cnt[t] = cnt.get(t, 0) + 1
        return cnt

    ta, ti = tok_counts(applied), tok_counts(ignored)
    only_ign = sorted(((t, n) for t, n in ti.items() if n >= 2 and ta.get(t, 0) == 0), key=lambda x: -x[1])
    only_app = sorted(((t, n) for t, n in ta.items() if n >= 2 and ti.get(t, 0) == 0), key=lambda x: -x[1])
    print("\nTitle words that appear ONLY in ignored roles (>=2x): " +
          (", ".join(f"{t} ({n})" for t, n in only_ign[:12]) or "-"))
    print("Title words that appear ONLY in applied roles (>=2x): " +
          (", ".join(f"{t} ({n})" for t, n in only_app[:12]) or "-"))

    def breakdown(field, fn=lambda v: v):
        rows: dict = {}
        for j in applied + ignored:
            k = fn(j.get(field) or "-")
            a_i = rows.setdefault(k, [0, 0])
            a_i[0 if j["status"] == "Applied" else 1] += 1
        return sorted(rows.items(), key=lambda kv: -(kv[1][0] + kv[1][1]))

    for field, label, fn in (("company", "Companies", lambda v: v), ("location", "Locations", norm_city),
                             ("source", "Sources", lambda v: v)):
        rows = [(k, v) for k, v in breakdown(field, fn) if v[0] + v[1] >= (2 if field == "company" else 1)]
        if rows:
            print(f"\n{label} (applied / ignored):")
            for k, (ap, ig) in rows[:10]:
                print(f"  {k}: {ap} / {ig}")

    locs = {norm_city(x) for x in read_lines(m / F_LOCATIONS)}
    outside = [j for j in applied if norm_city(j.get("location", "")) not in locs and "remote" not in locs]
    if outside:
        print("\nApplied roles OUTSIDE your saved locations: " +
              ", ".join(f"{j['id']} {j.get('location')}" for j in outside))
    hi_ign = [j for j in ignored if (j.get("score") or 0) >= max(75, threshold + 15)]
    lo_app = [j for j in applied if j.get("score") is not None and j["score"] < threshold + 10]
    if hi_ign:
        print("\nHigh-scoring roles you IGNORED (Laika may be over-rating something):")
        for j in hi_ign:
            print(f"  {j['id']} {j['score']} {j['title']} - {j['company']} | reason: {j.get('reason') or '-'}")
    if lo_app:
        print("\nLow-scoring roles you APPLIED to (Laika may be under-rating something):")
        for j in lo_app:
            print(f"  {j['id']} {j['score']} {j['title']} - {j['company']} | note: {j.get('reason') or '-'}")
    reasons = [j for j in sorted(applied + ignored, key=lambda j: j.get("status_date", ""), reverse=True) if j.get("reason")]
    if reasons:
        print("\nYour decision notes (newest first):")
        for j in reasons[:25]:
            print(f"  [{j['status']}] {j['id']} {j['title']} - {j['company']}: {j['reason']}")
    recent = sorted(applied + ignored, key=lambda j: j.get("status_date", ""), reverse=True)[:15]
    if recent:
        print("\nMost recent decisions:")
        for j in recent:
            print(f"  {j.get('status_date')} {j['status']:<8} {j.get('score') if j.get('score') is not None else '--':>3} "
                  f"{j['title']} - {j['company']} ({j.get('location') or '-'})")
    print("\nCurrent criteria: " + " | ".join(f"{k}={v}" for k, v in crit.items() if v))
    print("Roles: " + "; ".join(read_lines(m / F_ROLES)))
    print("Locations: " + "; ".join(read_lines(m / F_LOCATIONS)))
    if a.mark_reviewed:
        st["decisions_since_review"] = 0
        st["last_review_at"] = now_local().isoformat()
        save_state(m, st)
        print("\n(Marked as reviewed.)")


def cmd_rescore(a):
    """Re-score already registered roles (e.g. after the criteria changed). Decided roles keep their status."""
    m = get_master()
    st, db, crit = load_state(m), load_db(m), read_criteria(m)
    sync(m, db, st)
    raw = read_json(Path(a.input), None)
    if not isinstance(raw, list):
        die(f"Expected a JSON list in {a.input}: [{{\"id\": \"J0012\", \"score\": 55, \"why\": \"...\", \"concerns\": \"...\"}}]")
    threshold = crit_int(crit, "min_relevancy_score")
    changes = []
    for s in raw:
        j = find_job(db, str(s.get("id", "")))
        if not j:
            die(f"Unknown job id: {s.get('id')}")
        old_score, old_status = j.get("score"), j["status"]
        j["score"] = max(0, min(100, int(round(float(s["score"])))))
        for k in ("why", "concerns"):
            if s.get(k) is not None:
                j[k] = s[k].strip()
        if j["status"] in ("Pending", "Low Match"):
            j["status"] = "Pending" if j["score"] >= threshold else "Low Match"
            if j["status"] != old_status:
                j["status_date"] = today_str()
        write_card(m, j)
        changes.append((j, old_score, old_status))
    excel_ok = excel_write(m, db, st)
    save_db(m, db)
    save_state(m, st)
    for j, old_score, old_status in changes:
        moved = f" | {old_status} -> {j['status']}" if old_status != j["status"] else ""
        print(f"{j['id']} {j['title']} - {j['company']}: {old_score} -> {j['score']}{moved}")
    print(f"Re-scored {len(changes)} roles. Excel: {'updated' if excel_ok else 'NOT updated (file open)'}. "
          f"Still pending: {sum(1 for j in db['jobs'] if j['status'] == 'Pending')}")


def cmd_profile(a):
    if a.action == "here":
        target = (Path(a.workspace) if a.workspace and Path(a.workspace).is_dir() else Path.cwd()) / PROFILE_FILE
        if target.exists():
            print(f"This folder already has its own Laika profile: {target}")
        else:
            write_json(target, {})
            print(f"Created a separate Laika profile for this folder: {target}")
            print("Laika in this folder (and its subfolders) now starts fresh with onboarding; other folders keep "
                  "using their own profile. Delete the file to go back.")
        return
    cfg = load_config()
    print(f"Profile in use: {CONFIG_KIND} ({CONFIG_PATH})")
    print(f"Master folder: {cfg.get('master') or 'not set up yet'}")
    if CONFIG_KIND != "global":
        print(f"(The global profile {GLOBAL_CONFIG} is not used here.)")


def cmd_excel(a):
    m = get_master()
    st, db = load_state(m), load_db(m)
    sync(m, db, st)
    ok = excel_write(m, db, st)
    save_db(m, db)
    save_state(m, st)
    print(f"Excel {'written' if ok else 'NOT written'}: {m / F_TRACKER} ({len(db['jobs'])} roles)")


def cmd_runs(a):
    m = get_master()
    st = load_state(m)
    log = st.get("runs_log", [])
    if not log:
        print("No searches yet.")
        return
    for e in log[-a.n:]:
        print(f"{e['run']}: since {e.get('since')} | {e.get('sources')} | seen {e.get('seen')} | new {e.get('new')} | "
              f"shortlisted {e.get('shortlisted')} | duplicates {e.get('duplicates')}")


def cmd_doctor(a):
    print(f"Laika helper {VERSION} | Python {sys.version.split()[0]} at {sys.executable}")
    if sys.version_info < (3, 10):
        print("  PROBLEM: Laika needs Python 3.10 or newer.")
    req = Path(__file__).with_name("requirements.txt")
    missing = []
    for mod, why in (("openpyxl", "Excel tracker (required)"), ("pypdf", "PDF CV text extraction (optional)")):
        try:
            __import__(mod)
            print(f"  ok   {mod} - {why}")
        except ImportError:
            print(f"  MISSING {mod} - {why}")
            missing.append(mod)
    try:
        __import__("docx")
        print("  ok   python-docx - tailored CVs and cover letters (optional)")
    except ImportError:
        print("  MISSING python-docx - tailored CVs and cover letters (optional)")
        missing.append("python-docx")
    engines = load_cv_builder().pdf_engines()
    names = {"libreoffice": "LibreOffice", "word": "Microsoft Word"}
    print(f"  {'ok  ' if engines else 'NONE'} PDF converter for CVs and letters: "
          + (", ".join(names[e] for e in engines) if engines else "install LibreOffice (or use Microsoft Word on "
             "Windows) - without one, CVs and letters are Word files only"))
    if missing:
        print(f"  Install with:  \"{sys.executable}\" -m pip install -r \"{req}\"")
    print(f"Config: {CONFIG_PATH} ({'found' if CONFIG_PATH.exists() else 'not created yet'})")
    m = get_master(required=False)
    print(f"Master folder: {m or 'not set up'}")
    if a.net:
        try:
            page = http_get("https://www.linkedin.com/jobs-guest/jobs/api/seeMoreJobPostings/search?"
                            "keywords=engineer&location=London&f_TPR=r86400&start=0", tries=1)
            print(f"  LinkedIn guest search: {'ok' if parse_linkedin_cards(page) else 'reachable but no cards parsed'}")
        except FetchError as e:
            print(f"  LinkedIn guest search: FAILED ({e})")


# ----------------------------------------------------------------------------
# CLI
# ----------------------------------------------------------------------------
def build_parser() -> argparse.ArgumentParser:
    p = argparse.ArgumentParser(prog="laika.py", description="Laika, The Job Retriever - helper CLI")
    p.add_argument("--workspace", help="The folder Laika is used from (decides which profile applies)")
    sub = p.add_subparsers(dest="cmd", required=True)

    s = sub.add_parser("profile", help="Show which profile is in use, or `profile here` to give this folder its own")
    s.add_argument("action", nargs="?", choices=["show", "here"], default="show")
    s.set_defaults(func=cmd_profile)

    s = sub.add_parser("status", help="Overview: setup state, profile changes, last search, tracker counts")
    s.add_argument("-v", "--verbose", action="store_true", help="Also print the Laika CV summary")
    s.set_defaults(func=cmd_status)

    s = sub.add_parser("init", help="Create (or re-use) a master folder and remember it")
    s.add_argument("--master", required=True)
    s.add_argument("--cv", help="Path to the CV file to copy in")
    s.add_argument("--roles", nargs="*", help="Initial roles")
    s.add_argument("--locations", nargs="*", help="Initial locations")
    s.add_argument("--replace-profile", action="store_true",
                   help="For a folder that already has a Laika setup: replace its CV/roles/locations")
    s.set_defaults(func=cmd_init)

    s = sub.add_parser("doctor", help="Check dependencies and connectivity")
    s.add_argument("--net", action="store_true", help="Also test LinkedIn access")
    s.set_defaults(func=cmd_doctor)

    s = sub.add_parser("snapshot", help="Accept the current profile files as the baseline for change detection")
    s.set_defaults(func=cmd_snapshot)

    s = sub.add_parser("cv", help="Manage the CV: add FILE | text | path")
    s.add_argument("action", choices=["add", "text", "path"])
    s.add_argument("file", nargs="?")
    s.set_defaults(func=cmd_cv)

    s = sub.add_parser("certs", help="Licenses & certifications documents (optional): list | text | add FILE...")
    s.add_argument("action", choices=["list", "text", "add"])
    s.add_argument("files", nargs="*")
    s.set_defaults(func=cmd_certs)

    s = sub.add_parser("tailor", help="Prepare a tailored CV for one role (list number or job ID)")
    s.add_argument("ref")
    s.set_defaults(func=cmd_tailor)

    s = sub.add_parser("cv-build", help="Build a CV .docx + .pdf from content.json (or its folder) and check the rules")
    s.add_argument("target", help="A tailored-CV folder or a content.json file")
    s.add_argument("--name", help="Output file name without extension (default First_Last_CV_<Mon><YY>)")
    s.add_argument("--pdf", choices=["auto", "libreoffice", "word", "none"], default="auto",
                   help="PDF converter (auto: LibreOffice, else Microsoft Word)")
    s.set_defaults(func=cmd_cv_build)

    s = sub.add_parser("letter", help="Prepare a cover letter for one role (list number or job ID)")
    s.add_argument("ref")
    s.set_defaults(func=cmd_letter)

    s = sub.add_parser("letter-build", help="Build the cover letter .docx + .pdf from letter.json (or its folder)")
    s.add_argument("target", help="A cover-letter folder or a letter.json file")
    s.add_argument("--name", help="Output file name without extension (default First_Last_Cover_Letter_<Mon><YY>)")
    s.add_argument("--pdf", choices=["auto", "libreoffice", "word", "none"], default="auto")
    s.set_defaults(func=cmd_letter_build)

    s = sub.add_parser("style", help="Writing style for cover letters: show | path | add FILE... | samples")
    s.add_argument("action", choices=["show", "path", "add", "samples"])
    s.add_argument("files", nargs="*")
    s.add_argument("--chars", type=int, default=6000, help="Max characters printed per sample")
    s.set_defaults(func=cmd_style)

    s = sub.add_parser("list", help="Show or edit roles / locations / sites")
    s.add_argument("action", choices=["show", "set", "add", "remove"])
    s.add_argument("kind", choices=list(LIST_SPECS))
    s.add_argument("items", nargs="*")
    s.add_argument("--reason", default="")
    s.set_defaults(func=cmd_list)

    s = sub.add_parser("criteria", help="Show or set search criteria")
    s.add_argument("action", choices=["show", "set"])
    s.add_argument("key", nargs="?")
    s.add_argument("value", nargs="?", default="")
    s.add_argument("--reason", default="")
    s.set_defaults(func=cmd_criteria)

    s = sub.add_parser("run", help="Search runs: start [--new] | status | done QIDS | abandon")
    s.add_argument("action", choices=["start", "status", "done", "abandon"])
    s.add_argument("qids", nargs="*")
    s.add_argument("--new", action="store_true", help="Discard an open run and start fresh")
    s.set_defaults(func=cmd_run)

    s = sub.add_parser("fetch-linkedin", help="Fetch LinkedIn results for the open run")
    s.add_argument("--max-per-query", type=int, default=0)
    s.add_argument("--delay", type=float, default=1.5, help="Seconds between requests (randomised)")
    s.add_argument("--queries", help="Only these query ids, comma separated")
    s.add_argument("--refetch", action="store_true")
    s.set_defaults(func=cmd_fetch_linkedin)

    s = sub.add_parser("add-candidates", help="Add jobs found by browsing / web search (JSON file)")
    s.add_argument("--in", dest="input", required=True)
    s.add_argument("--source", required=True, help="Site name, e.g. kariyer.net")
    s.add_argument("--query", help="Plan query ids these results belong to (marks them done)")
    s.add_argument("--no-date-filter", action="store_true", help="Do not drop postings older than the window")
    s.set_defaults(func=cmd_add_candidates)

    s = sub.add_parser("details", help="Fetch LinkedIn job descriptions for candidates")
    s.add_argument("--limit", type=int, default=40)
    s.add_argument("--cids", help="Specific candidate ids, comma separated")
    s.add_argument("--min-title-match", type=float, default=0.01)
    s.add_argument("--delay", type=float, default=1.2)
    s.set_defaults(func=cmd_details)

    s = sub.add_parser("candidates", help="List the open run's candidates; --export writes the scoring file")
    s.add_argument("--export", action="store_true")
    s.add_argument("--all", action="store_true", help="Export flagged / unrelated titles too")
    s.add_argument("--chars", type=int, default=2500, help="Description length in the export")
    s.add_argument("--cids", help="Export only these candidate ids (comma separated)")
    s.add_argument("--excerpt", action="store_true",
                   help="Export a compact excerpt (opening, remote/location mentions, requirements) per role")
    s.add_argument("--compact", action="store_true", help="List view: one short line per candidate, by location")
    s.add_argument("--out", help="Write the list view to this file instead of the console")
    s.set_defaults(func=cmd_candidates)

    s = sub.add_parser("register", help="Register the scored run: database, Excel, report, shortlist")
    s.add_argument("--scores", help="Scores JSON (default: the run folder's scores.json)")
    s.add_argument("--skipped", help="Query ids that could not be searched (they keep their old window)")
    s.set_defaults(func=cmd_register)

    s = sub.add_parser("decide", help="Record decisions, e.g. --apply 1,3 --ignore 2 --reason 2=too junior")
    s.add_argument("--apply", default="")
    s.add_argument("--ignore", default="")
    s.add_argument("--pending", default="")
    s.add_argument("--reason", action="append", help="REF=text (repeatable)")
    s.set_defaults(func=cmd_decide)

    s = sub.add_parser("pending", help="List roles waiting for a decision")
    s.add_argument("-v", "--verbose", action="store_true")
    s.set_defaults(func=cmd_pending)

    s = sub.add_parser("show", help="Show full details for list numbers or job IDs")
    s.add_argument("ref")
    s.set_defaults(func=cmd_show)

    s = sub.add_parser("feedback", help="Decision statistics for reviewing the search criteria")
    s.add_argument("--mark-reviewed", action="store_true")
    s.set_defaults(func=cmd_feedback)

    s = sub.add_parser("rescore", help="Re-score registered roles after a criteria change (JSON list of id/score/why/concerns)")
    s.add_argument("--in", dest="input", required=True)
    s.set_defaults(func=cmd_rescore)

    s = sub.add_parser("excel", help="Re-sync and rewrite the Excel tracker")
    s.set_defaults(func=cmd_excel)

    s = sub.add_parser("runs", help="Show recent searches")
    s.add_argument("-n", type=int, default=10)
    s.set_defaults(func=cmd_runs)
    return p


def main(argv=None):
    global CONFIG_PATH, CONFIG_KIND
    args = build_parser().parse_args(argv)
    if args.workspace:
        CONFIG_PATH, CONFIG_KIND = resolve_config(args.workspace)
    if args.cmd == "cv" and args.action == "add" and not args.file:
        die("cv add needs a file path")
    if args.cmd == "criteria" and args.action == "set" and not args.key:
        die("criteria set needs KEY VALUE")
    args.func(args)


if __name__ == "__main__":
    main()
